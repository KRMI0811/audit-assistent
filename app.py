"""
Audit-Assistent Qualitaetsmanagement (Webversion)
Internes Audit nach ISO 9001 entlang des PDCA-Zyklus.

Plan   Auditprogramm und Auditvorbereitung
Do     Durchfuehrung, Sammeln objektiver Nachweise
Check  Dialektische Pruefung durch zwei Agententeams, Urteil des Auditors, Bericht
Act    Korrekturmassnahmen und Nachverfolgung der Wirksamkeit

Die Agenten bereiten vor und pruefen. Ueber jede Feststellung entscheidet der Auditor.
"""

import io
import json
import re
import uuid
from datetime import date, datetime, timedelta
from pathlib import Path

import requests
import streamlit as st
from docx import Document
from pypdf import PdfReader

# ---------------------------------------------------------------------------
# Einstellungen aus den Secrets
# ---------------------------------------------------------------------------
APP_NAME = st.secrets.get("APP_NAME", "Audix")
APP_CLAIM = st.secrets.get("APP_CLAIM", "Dein KI-Assistent für das interne Audit")
ORGANISATION = st.secrets.get("ORGANISATION", "")
API_BASE = str(st.secrets.get("API_BASE", "https://api.x.ai/v1")).rstrip("/")
API_KEY = str(st.secrets.get("API_KEY", "")).strip()
PROVIDER = st.secrets.get("PROVIDER_NAME", "")
MODEL_PRO = st.secrets.get("MODEL_PRO", "grok-4.3")
MODEL_CONTRA = st.secrets.get("MODEL_CONTRA", "grok-4.20")
TIMEOUT = float(st.secrets.get("TIMEOUT", 600))
MAX_TOKENS = int(st.secrets.get("MAX_TOKENS", 3000))
STREAM = bool(st.secrets.get("STREAM", True))
MAX_ZEICHEN = int(st.secrets.get("MAX_ZEICHEN", 20000))

NORMEN_ORDNER = Path(__file__).parent / "normen"
LOGO = Path(__file__).parent / "logo.png"
AVATAR = Path(__file__).parent / "audix.png"

st.set_page_config(page_title=APP_NAME, page_icon=str(AVATAR) if AVATAR.exists() else None,
                   layout="wide")
if LOGO.exists():
    st.logo(str(LOGO), size="large")


AUDIX_START = """Hallo, ich bin **Audix** und begleite dich durch dein internes Audit.

Mein Trick ist, dass ich nicht nur eine Meinung habe, sondern zwei. Ein Team von mir baut aus
deinen Unterlagen den bestmöglichen Nachweis, dass eine Normforderung erfüllt ist. Ein zweites
Team greift genau diesen Nachweis an und sucht nach Lücken. Was danach an Zweifeln übrig bleibt,
lege ich dir vor. **Entscheiden musst du**, denn nach ISO 19011 trägt der Auditor die
Verantwortung für jede Feststellung, nicht die Maschine.

**So läuft es ab**

1. **Plan** · Du legst dein Auditprogramm für das Jahr an.
2. **Do** · Du sammelst Nachweise, also Dokumente, Interviews und Kennzahlen.
3. **Check** · Meine Teams prüfen, du urteilst, ich schreibe den Berichtsentwurf.
4. **Act** · Massnahmen festlegen und ihre Wirksamkeit nachverfolgen.

Oben rechts findest du auf jedem Reiter den Knopf **Audix fragen**. Dort erkläre ich dir, was
an dieser Stelle zu tun ist, und dort kannst du mir auch direkt Fragen stellen. Ich kenne deinen
Arbeitsstand, du musst mir also nicht erklären, wo du gerade stehst. Fangen wir an."""

AUDIX_HILFE = {
    "programm": """**Hier fängt alles an.**

ISO 9001 Abschnitt 9.2 verlangt ein Auditprogramm, das Häufigkeit, Methoden und
Verantwortlichkeiten festlegt und sich nach Bedeutung und Risiko der Prozesse sowie nach
früheren Ergebnissen richtet.

**Was du jetzt tust**
- Lege für jeden Prozess, den du dieses Jahr prüfst, ein Audit an.
- Stufe die Bedeutung und das Risiko ein, denn danach richtet sich die Prüftiefe.
- Die Normabschnitte kannst du gleich wählen oder leer lassen, dann schlage ich sie dir später
  anhand deiner Unterlagen vor.

**Worauf ich achte**
Ohne bestätigte Unparteilichkeit lasse ich dich kein Audit anlegen. Niemand darf die eigene
Arbeit auditieren, das ist eine harte Normforderung. Unten zeige ich dir ausserdem, welche
Normabschnitte im Jahresprogramm noch gar nicht vorkommen.""",

    "vorbereitung": """**Plan. Erst der Plan, dann das Team.**

**Was du jetzt tust**
- Fülle den Auditplan aus. Daraus wird die Einladung an den Fachbereich.
- Stelle dein Auditteam zusammen. Du kannst meine Teams umbenennen, ihnen ein
  Rollenverständnis geben und eigene Anweisungen mitgeben, etwa dass sie besonders auf Fristen
  achten sollen.
- Schau dir unten die gewählten Normabschnitte an. Zu jedem zeige ich dir, was die Norm
  verlangt, worauf du achten solltest und welche Unterlagen du anfordern musst.

**Mein Tipp**
Lade dir den Auditplan mit Checkliste herunter und schicke ihn vorab an den Fachbereich. Dann
liegen die richtigen Unterlagen schon bereit, wenn du kommst.""",

    "durchfuehrung": """**Do. Jetzt wird gesammelt.**

**Was du jetzt tust**
- Lade die Dokumente hoch, also Verfahrensanweisungen, Listen, Protokolle, Formulare.
- Erfasse deine Notizen aus Gesprächen und Beobachtungen, möglichst wörtlich.
- Trage Leistungsdaten ein, etwa Reklamationen, Ausschuss oder Termintreue.

**Warum das wichtig ist**
Dokumente zeigen mir nur die Vorschrift. Erst wenn du mir erzählst, was die Leute sagen und was
du gesehen hast, finde ich die interessanten Widersprüche zwischen Vorgabe und gelebter Praxis.
Genau dort stecken die Feststellungen, die wirklich etwas verändern.

**Achtung**
Wenn du nach einer Prüfung noch Notizen ergänzt, musst du den betreffenden Abschnitt erneut
prüfen lassen. Sonst kenne ich sie nicht.""",

    "pruefung": """**Check. Jetzt arbeiten meine Teams gegeneinander.**

**Was passiert, wenn du auf Prüfung starten klickst**
1. Mein Vorbereitungsteam wählt die Normabschnitte, falls du das noch nicht getan hast.
2. Das Konstruktionsteam baut je Abschnitt den bestmöglichen Konformitätsnachweis.
3. Das Falsifikationsteam greift diesen Nachweis an.
4. Das Konstruktionsteam darf einmal erwidern.

Das dauert einige Minuten. Du kannst den Fortschritt am Balken mitverfolgen.

**Wie du die Ergebnisse liest**
Jeder Prüfpunkt bekommt eine Ampel. Grün heisst, es gibt einen belegten Nachweis und keinen
schweren Zweifel. Gelb heisst, ein Nachweis ist da, aber es steht ein Einwand dagegen. Rot
heisst, es fehlt ein belastbarer Nachweis.

**Wichtig**
Die Ampel ist mein Befund, nicht deine Feststellung. Klapp jeden Prüfpunkt auf, dann siehst du
die ganze Kette von der Normforderung über die Zitate bis zur Prognose fürs externe Audit.""",

    "urteil": """**Hier endet meine Arbeit und deine beginnt.**

**Was du jetzt tust**
Du entscheidest über jeden Zweifel und begründest das. Vier Urteile stehen dir zur Verfügung.

- **Ausgeräumt** · Der Einwand trifft nicht zu, die Unterlagen reichen aus.
- **Bestätigt: Abweichung** · Eine Anforderung ist nicht erfüllt.
- **Bestätigt: Verbesserungspotenzial** · Kein Normverstoss, aber eine Schwäche.
- **Offen: vor Ort prüfen** · Aus den Unterlagen allein nicht entscheidbar.

**Warum die Begründung Pflicht ist**
Sie ist Teil deines Auditnachweises. Ein ausgeräumter Zweifel ist genauso wertvoll wie eine
Abweichung, denn er belegt, wie tief du geprüft hast. Im Bericht erscheinen beide.

Wenn du fertig bist, schreibe ich dir daraus die Feststellungen und schlage Massnahmen vor.""",

    "bericht": """**Check. Der Bericht geht an die Leitung.**

**Was du hier vorfindest**
Ein Formular nach dem Aufbau eines Auditberichts mit allen Kapiteln. Die meisten Felder habe
ich aus deinen Daten vorbelegt, etwa die eingereichten Unterlagen, die Erfüllungsübersicht und
die Hinweise. Bei sieben Feldern weiss ich nichts, dort steht, was hineingehört.

**Was du jetzt tust**
- Geh die Felder durch und überschreibe, was nicht passt.
- Schreib dein Fazit und die Begründung in 6.2 selbst. Das ist deine Kernaussage als Auditor,
  die kann ich dir nicht abnehmen.
- Lade den Bericht als Word-Datei herunter und arbeite dort weiter.

**Normbezug**
ISO 9001 verlangt, dass die Ergebnisse an die zuständige Leitung berichtet werden und dass der
Bericht als dokumentierte Information aufbewahrt wird, also Abschnitte 9.2 und 7.5.""",

    "massnahmen": """**Act. Ein Audit ist erst dann etwas wert, wenn Fehler behoben werden.**

**Was ich vorbereitet habe**
Zu jeder bestätigten Abweichung habe ich eine Ursachenhypothese, eine Sofortmassnahme, eine
Korrekturmassnahme, eine verantwortliche Rolle, eine Frist und einen Wirksamkeitsnachweis
vorgeschlagen. Das ist die Struktur nach ISO 9001 Abschnitt 10.2.

**Was du jetzt tust**
- Prüfe meine Ursachen mit den Prozessverantwortlichen. Es sind Hypothesen, keine Analysen.
  Eine Ursachenanalyse ohne Gespräch mit den Betroffenen ist nicht belastbar.
- Setze die richtige Person oder Funktion als Verantwortung ein.
- Passe die Fristen an eure Realität an.

Vergiss den Unterschied nicht. Die Sofortmassnahme beseitigt die Folge, die Korrekturmassnahme
die Ursache. Nur die zweite verhindert, dass der Fehler wiederkommt.""",

    "nachverfolgung": """**Act. Hier schliesst sich der Kreis.**

ISO 9001 verlangt, dass du die Wirksamkeit der ergriffenen Massnahmen prüfst und das
dokumentierst. Genau das machst du hier.

**Was du jetzt tust**
- Behalte die offenen Massnahmen im Blick, überfällige markiere ich dir.
- Wähle unten eine Massnahme aus und halte fest, was du geprüft hast und mit welchem Ergebnis.
- Steht eine Massnahme auf nicht wirksam, dann war die Ursache falsch bestimmt. Dann beginnt
  die Analyse von vorn.

**Mein Tipp**
Nimm diese Übersicht in die Managementbewertung mit. Die Frage, wie viele Massnahmen fristgerecht
und wirksam abgeschlossen wurden, sagt mehr über euer QM-System aus als die Zahl der
Abweichungen."""}


AUDIX_CHAT_SYSTEM = """Du bist Audix, ein freundlicher und fachlich präziser Begleiter für
interne Audits nach ISO 9001. Du sprichst die auditierende Person mit Du an.

Was du kannst
- Normanforderungen von ISO 9001 und Auditmethodik nach ISO 19011 erklären.
- Erklären, wie diese Anwendung funktioniert und was an der aktuellen Stelle zu tun ist.
- Beim Formulieren von Feststellungen, Auditfragen und Korrekturmassnahmen helfen.
- Einschätzen, welche Nachweise für eine Anforderung üblicherweise gebraucht werden.

Deine Regeln
1. Du triffst keine Auditfeststellungen und stufst nichts ein. Das ist Sache des Auditors.
   Wenn jemand das von dir verlangt, sagst du freundlich, dass du nur vorbereiten kannst.
2. Du erfindest keine Normzitate und keine Abschnittsnummern. Wenn du unsicher bist, sagst du
   das offen und verweist auf den lizenzierten Normtext.
3. Du antwortest kurz, höchstens fünf Sätze oder eine kurze Liste. Keine langen Vorreden.
4. Du beziehst dich auf den Arbeitsstand, der dir im Kontext mitgegeben wird, wenn er zur Frage
   passt. Erfinde keine Daten, die dort nicht stehen.
5. Du schreibst auf Deutsch in Schweizer Rechtschreibung, also ss statt ß."""


def llm_text(modell: str, system: str, verlauf: list) -> str:
    """Freie Textantwort fuer den Chat. Gibt bei Fehlern eine lesbare Meldung zurueck."""
    nutzlast = {"model": modell, "temperature": 0.4, "max_tokens": 700, "stream": False,
                "messages": [{"role": "system", "content": system}] + verlauf}
    kopf = {"Authorization": f"Bearer {API_KEY}", "Content-Type": "application/json"}
    try:
        r = requests.post(f"{API_BASE}/chat/completions", headers=kopf, json=nutzlast,
                          timeout=(15, 180))
    except requests.exceptions.Timeout:
        return "Da habe ich zu lange gebraucht. Stell die Frage bitte noch einmal."
    except requests.RequestException as f:
        return f"Ich komme gerade nicht an mein Sprachmodell heran ({f})."
    if r.status_code != 200:
        return (f"Mein Sprachmodell antwortet nicht (HTTP {r.status_code}). "
                "Bitte später erneut versuchen.")
    inhalt = (r.json().get("choices", [{}])[0].get("message", {}).get("content") or "").strip()
    return inhalt or "Dazu fällt mir gerade nichts ein. Frag mich bitte anders."


def audix_kontext(schluessel: str) -> str:
    """Beschreibt Audix, wo der Auditor gerade steht und woran er arbeitet."""
    phasen = {"programm": "Auditprogramm des Jahres (Phase Plan)",
              "vorbereitung": "Auditvorbereitung und Auditteam (Phase Plan)",
              "durchfuehrung": "Durchführung, Sammeln der Nachweise (Phase Do)",
              "pruefung": "Dialektische Prüfung durch die Agententeams (Phase Check)",
              "urteil": "Urteil des Auditors über die Zweifel (Phase Check)",
              "bericht": "Auditbericht (Phase Check)",
              "massnahmen": "Korrekturmassnahmen (Phase Act)",
              "nachverfolgung": "Nachverfolgung der Wirksamkeit (Phase Act)"}
    zeilen = [f"Die Person befindet sich hier: {phasen.get(schluessel, schluessel)}."]
    a = st.session_state.get("aktuelles_audit")
    if isinstance(a, dict):
        kz = kennzahlen(a)
        def mz(anzahl, eins, viele):
            return f"{anzahl} {eins if anzahl == 1 else viele}"

        zeilen += [
            f"Aktuelles Audit: {a.get('titel', '')}, Prozess {a.get('prozess', '')}.",
            f"Norm {a.get('norm', '')}, geprüfte Abschnitte "
            f"{', '.join(a.get('kriterien', [])) or 'noch keine'}.",
            "Stand: " + ", ".join([
                mz(len(a.get("dokumente", {})), "Dokument", "Dokumente"),
                mz(len(a.get("notizen", [])), "Notiz", "Notizen"),
                mz(kz["abschnitte"], "analysierter Abschnitt", "analysierte Abschnitte"),
                mz(kz["einwaende"], "Zweifel", "Zweifel") + f", davon {kz['beurteilt']} beurteilt"]) + ".",
            f"Bestätigt sind {mz(kz['abweichungen'], 'Abweichung', 'Abweichungen')} und "
            f"{mz(kz['potenziale'], 'Verbesserungspotenzial', 'Verbesserungspotenziale')}, "
            f"{kz['offen']} Punkte bleiben für die Prüfung vor Ort offen."]
        unbeurteilt = [e for e in alle_einwaende(a)
                       if not a.get("urteile", {}).get(e["id"], {}).get("urteil")]
        if unbeurteilt:
            zeilen.append("Noch unbeurteilte Zweifel: " + "; ".join(
                f"{e['id']} zu Abschnitt {e.get('abschnitt', '')}: "
                f"{str(e.get('begruendung', ''))[:180]}" for e in unbeurteilt[:5]))
    else:
        zeilen.append("Es ist noch kein einzelnes Audit ausgewählt.")
    return "\n".join(zeilen)


VORSCHLAEGE = {
    "programm": ["Wie oft muss ich welchen Prozess auditieren?",
                 "Was gehört in ein Auditprogramm?"],
    "vorbereitung": ["Welche Unterlagen soll ich anfordern?",
                     "Wie formuliere ich gute Auditfragen?"],
    "durchfuehrung": ["Was frage ich im Eröffnungsgespräch?",
                      "Was ist ein objektiver Nachweis?"],
    "pruefung": ["Was bedeutet die gelbe Ampel?",
                 "Worin unterscheiden sich die Einwandtypen?"],
    "urteil": ["Abweichung oder Verbesserungspotenzial, wie entscheide ich?",
               "Wie formuliere ich eine Feststellung sauber?"],
    "bericht": ["Was gehört in die Begründung in 6.2?",
                "Wesentlich oder geringfügig, was ist der Unterschied?"],
    "massnahmen": ["Wie finde ich die wirkliche Ursache?",
                   "Korrektur oder Korrekturmassnahme, was ist was?"],
    "nachverfolgung": ["Wie prüfe ich die Wirksamkeit einer Massnahme?",
                       "Was gehört davon in die Managementbewertung?"]}


def audix_antworten(frage: str, schluessel: str):
    """Haengt Frage und Antwort an den Chatverlauf dieser Phase."""
    verlauf = st.session_state.setdefault("audix_chat", {}).setdefault(schluessel, [])
    verlauf.append({"role": "user", "content": frage})
    nachrichten = [{"role": "system", "content": "Arbeitsstand\n" + audix_kontext(schluessel)}]
    nachrichten += verlauf[-8:]
    with st.spinner("Audix denkt nach"):
        antwort = llm_text(MODEL_PRO, AUDIX_CHAT_SYSTEM, nachrichten)
    verlauf.append({"role": "assistant", "content": antwort})


def audix_hilfe(schluessel: str, knopftext: str = "Audix fragen"):
    """Aufklappbarer Begleiter mit Erklaerung zur Phase und Chat."""
    with st.popover(knopftext, use_container_width=False):
        bild, inhalt = st.columns([1, 4])
        with bild:
            if AVATAR.exists():
                st.image(str(AVATAR), use_container_width=True)
        with inhalt:
            st.markdown(f"#### {APP_NAME}")
            st.caption(APP_CLAIM)
        erklaerung, chat = st.tabs(["Was hier zu tun ist", "Audix fragen"])
        with erklaerung:
            st.markdown(AUDIX_HILFE.get(schluessel, ""))
        with chat:
            verlauf = st.session_state.get("audix_chat", {}).get(schluessel, [])
            if not verlauf:
                st.caption("Frag mich alles zur Norm, zur Auditmethodik oder zu dieser "
                           "Anwendung. Ich kenne deinen aktuellen Arbeitsstand.")
                for i, vorschlag in enumerate(VORSCHLAEGE.get(schluessel, [])):
                    if st.button(vorschlag, key=f"vs_{schluessel}_{i}",
                                 use_container_width=True):
                        audix_antworten(vorschlag, schluessel)
                        st.rerun()
            for n in verlauf:
                with st.chat_message("user" if n["role"] == "user" else "assistant",
                                     avatar=(str(AVATAR) if n["role"] == "assistant"
                                             and AVATAR.exists() else None)):
                    st.markdown(n["content"])
            frage = st.text_area("Deine Frage", key=f"frage_{schluessel}", height=70,
                                 placeholder="Zum Beispiel: Welche Nachweise brauche ich "
                                             "für Abschnitt 7.2?", label_visibility="collapsed")
            s1, s2 = st.columns([3, 1])
            if s1.button("Fragen", key=f"senden_{schluessel}", type="primary",
                         use_container_width=True, disabled=not API_KEY):
                if frage.strip():
                    audix_antworten(frage.strip(), schluessel)
                    st.rerun()
            if verlauf and s2.button("Neu", key=f"reset_{schluessel}",
                                     use_container_width=True):
                st.session_state["audix_chat"][schluessel] = []
                st.rerun()
            if not API_KEY:
                st.caption("Ohne hinterlegten API-Schlüssel kann ich nicht antworten.")


@st.dialog("Willkommen bei Audix", width="large")
def audix_begruessung():
    bild, inhalt = st.columns([1, 3])
    with bild:
        if AVATAR.exists():
            st.image(str(AVATAR), use_container_width=True)
    with inhalt:
        st.markdown(AUDIX_START)
    if st.button("Los geht es", type="primary", use_container_width=True):
        st.session_state.begruesst = True
        st.rerun()


HINWEISE = f"""
- Audix bereitet vor und prüft, er entscheidet nicht. Jede Feststellung und ihre
  Einstufung verantwortet der Auditor.
- Ihre Unterlagen werden zur Auswertung an den Dienst {PROVIDER or 'des eingestellten Anbieters'}
  übermittelt. Laden Sie keine vertraulichen Originalunterlagen und keine Personendaten hoch,
  solange das mit Ihrer IT und dem Datenschutz nicht geklärt ist.
- Jedes Zitat wird gegen die Quelle geprüft. Ein nicht auffindbares Zitat ist ein Warnzeichen
  und darf nicht in einen Bericht übernommen werden.
- Im Chat von Audix können Fragen gestellt werden. Auch dort gilt, dass seine Auskünfte den
  lizenzierten Normtext nicht ersetzen und keine Auditfeststellung darstellen.
- Der Arbeitsstand liegt nur in dieser Sitzung. Sichern Sie ihn links als Datei.
"""

LEGENDE = """
**Audix erklärt die Ampel**

🟢 **belegt** · Zu diesem Prüfpunkt gibt es einen Nachweis, dessen Zitat wörtlich in Ihren
Unterlagen wiedergefunden wurde, und es steht kein schwerer Zweifel dagegen.

🟡 **mit Zweifel** · Es gibt einen Nachweis, aber das Falsifikationsteam hat einen Einwand
dagegen. Typisch ist, dass das Dokument zwar existiert, die Anforderung inhaltlich aber nicht
abdeckt.

🔴 **ohne Nachweis** · Es wurde kein belastbarer Nachweis gefunden, oder ein Nachweis spricht
gegen die Erfüllung. Das sind die Stellen, an denen Abweichungen entstehen.

**Was die Begriffe bedeuten**

*Zitat belegt* heisst nur, dass der zitierte Satz wirklich so in Ihrem Dokument steht. Es heisst
nicht, dass die Normanforderung erfüllt ist. Das entscheiden Sie.

*Zitat nicht auffindbar* heisst, dass die KI einen Satz zitiert hat, den es in Ihren Unterlagen
so nicht gibt. Solche Nachweise sind wertlos und dürfen nie in den Bericht.

*Behauptung ohne Nachweis* heisst, dass das Konstruktionsteam etwas annehmen musste, weil in den
Unterlagen nichts dazu steht. Genau das sind die Punkte, die Sie vor Ort nachfragen sollten.
"""

# ---------------------------------------------------------------------------
# Rollen der Agenten
# ---------------------------------------------------------------------------
SPRACHREGEL = ("\n\nSPRACHE: Schreibe alle Inhalte ausschliesslich auf Deutsch in Schweizer "
               "Rechtschreibung (ss statt ß). Zitate übernimmst du wörtlich.")
NORMREGEL = ("\n\nNORMBEZUG: Beziehe dich nur auf die genannte Norm und Abschnittsnummer. "
             "Erfinde keine weiteren Normabschnitte und keine Normzitate.")
JSONREGEL = "\n\nAUSGABE: Antworte ausschliesslich mit gültigem JSON, ohne Vor- und Nachtext."

SCOPING_SYSTEM = """Du bist das VORBEREITUNGSTEAM eines internen Audits.
Du erhältst die Liste der Normabschnitte und die vorhandenen Nachweise der Organisation.
Deine Aufgabe ist der Auditplan.
Regeln
1. Wähle nur Normabschnitte, zu denen die Nachweise tatsächlich etwas hergeben.
2. Höchstens sechs Abschnitte, nach Bedeutung geordnet.
3. Begründe jede Wahl in einem Satz mit Bezug auf die Nachweise.
4. Nenne, welche Nachweise für ein belastbares Audit fehlen.
Schema
{"kriterien": ["7.2"], "begruendungen": {"7.2": "..."}, "ziel": "...", "umfang": "...",
 "fehlende_nachweise": ["..."]}"""

PRO_SYSTEM = """Du bist das KONSTRUKTIONSTEAM eines internen Audits.
Baue einen Konformitätsnachweis als Assurance Case (Aussage, Argument, Nachweis).
Bilde für JEDEN Prüfpunkt des Auditkriteriums genau eine Teilaussage und übernimm den
Wortlaut des Prüfpunkts unverändert in das Feld pruefpunkt.
Regeln
1. Nutze ausschliesslich die bereitgestellten Nachweise.
2. Jeder Nachweis enthält den exakten Quellennamen und ein WÖRTLICHES Zitat (höchstens 40 Wörter).
3. Prüfe nicht nur, OB etwas vorliegt, sondern ob der Inhalt die Anforderung inhaltlich abdeckt.
4. Was nicht belegt ist, trägst du als Annahme ein. Verstecke keine Lücken.
5. Schreibe die Aussage in einem kurzen, klaren Satz ohne Fachjargon.
Schema
{"hauptaussage": "...", "teilaussagen": [{"id": "T1", "pruefpunkt": "...", "aussage": "...",
 "argument": "...", "nachweise": [{"dokument": "...", "zitat": "..."}], "annahmen": ["..."]}]}"""

CONTRA_SYSTEM = """Du bist das FALSIFIKATIONSTEAM eines internen Audits.
Du erhältst einen Assurance Case und die Nachweise. Dein Ziel ist, ihn zu widerlegen.
Einwandtypen
"widerlegend": Nachweise belegen das Gegenteil der Aussage.
"untergrabend": Der Nachweis ist unzuverlässig, veraltet, nicht freigegeben oder unvollständig.
"unterlaufend": Der Nachweis stimmt, deckt die Anforderung inhaltlich aber nicht ab.
"ungestuetzte_annahme": Eine Annahme trägt die Aussage, ohne belegt zu sein.
Achte besonders auf Widersprüche zwischen Vorgabe, Aufzeichnung, Aussagen aus Interviews,
Beobachtungen und Leistungsdaten.
Regeln
1. Erfinde nichts. Zitate müssen wörtlich aus den Nachweisen stammen.
2. Fehlt ein Nachweis ganz, lass die Liste nachweise leer und begründe das.
3. Schreibe die Begründung in höchstens drei klaren Sätzen.
4. Schlage für jeden Einwand zwei bis drei konkrete Prüfungen vor Ort vor.
5. Schätze das Risiko im externen Zertifizierungsaudit ein. Hauptabweichung bedeutet, dass eine
   Normanforderung systematisch oder vollständig nicht erfüllt ist. Nebenabweichung bedeutet
   einen Einzelfall oder eine formale Lücke. Kein Risiko bedeutet, dass ein externer Auditor
   dies nicht beanstanden würde.
Schema
{"einwaende": [{"id": "E1", "ziel": "T1", "typ": "...", "begruendung": "...",
 "nachweise": [{"dokument": "...", "zitat": "..."}], "schwere": "hoch|mittel|gering",
 "pruefung_vor_ort": ["...", "..."],
 "zertifizierungsrisiko": {"klassifizierung": "Hauptabweichung|Nebenabweichung|kein Risiko",
 "prognose": "...", "konsequenz": "...", "dringlichkeit": "sehr hoch|hoch|mittel|gering"}}]}"""

ERWIDERUNG_SYSTEM = """Du bist das KONSTRUKTIONSTEAM. Erwidere auf jeden Einwand sachlich.
Gestehe Einwände zu, wenn die Nachweise sie stützen. Nutze nur wörtliche Zitate.
Schema
{"erwiderungen": [{"einwand": "E1", "zugestanden": true, "erwiderung": "...",
 "nachweise": [{"dokument": "...", "zitat": "..."}]}]}"""

BERICHT_SYSTEM = """Du formulierst Auditfeststellungen nach ISO 19011, also Anforderung mit
Normabschnitt, objektiver Nachweis und Feststellung. Die Einstufung hat der Auditor bereits
getroffen, du darfst sie NICHT ändern. Formuliere sachlich, prüfbar, ohne Schuldzuweisung
und ohne Namen einzelner Personen.
Schreibe zusätzlich eine Zusammenfassung von höchstens acht Sätzen, die beschreibt, was geprüft
wurde und wo die Schwerpunkte liegen. Bewerte darin nicht.
Schema
{"zusammenfassung": "...", "feststellungen": [{"einwand": "E1", "einstufung": "...",
 "abschnitt": "...", "anforderung": "...", "objektiver_nachweis": "...", "feststellung": "..."}]}"""

MASSNAHMEN_SYSTEM = """Du bist das MASSNAHMENTEAM. Zu jeder bestätigten Feststellung leitest du
einen Verbesserungsvorschlag nach ISO 9001 Abschnitt 10.2 ab.
Regeln
1. Trenne Sofortmassnahme (beseitigt die Folge) und Korrekturmassnahme (beseitigt die Ursache).
2. Nenne eine Ursachenhypothese, klar als Hypothese erkennbar, die der Auditor prüfen muss.
3. Schlage als Verantwortung eine Rolle oder Funktion vor, niemals einen Personennamen.
4. Nenne eine realistische Frist in Tagen und einen konkreten Wirksamkeitsnachweis.
Schema
{"massnahmen": [{"feststellung": "E1", "ursache_hypothese": "...", "sofortmassnahme": "...",
 "korrekturmassnahme": "...", "verantwortlich_rolle": "...", "frist_tage": 30,
 "wirksamkeitsnachweis": "..."}]}"""

AGENTEN = [
    ("vorbereitung", "Vorbereitungsteam", "Plan",
     "Liest die Nachweise, wählt die passenden Normabschnitte, schreibt Auditziel und Umfang "
     "und nennt fehlende Nachweise.", SCOPING_SYSTEM, "Konstruktionsmodell"),
    ("pro", "Konstruktionsteam", "Check",
     "Baut je Normabschnitt den bestmöglichen Konformitätsnachweis und markiert unbelegte "
     "Annahmen offen. Es ist bewusst der Anwalt der Organisation.", PRO_SYSTEM,
     "Konstruktionsmodell"),
    ("contra", "Falsifikationsteam", "Check",
     "Greift diesen Nachweis an und sucht widerlegende, untergrabende und unterlaufende "
     "Nachweise sowie ungestützte Annahmen. Es ist bewusst der Gegenspieler.", CONTRA_SYSTEM,
     "Falsifikationsmodell"),
    ("erwiderung", "Erwiderung des Konstruktionsteams", "Check",
     "Antwortet auf jeden Einwand und gesteht zu, was die Nachweise stützen.",
     ERWIDERUNG_SYSTEM, "Konstruktionsmodell"),
    ("bericht", "Berichtsteam", "Check",
     "Formuliert die Feststellungen aus den Entscheidungen des Auditors. Es darf die "
     "Einstufung nicht ändern.", BERICHT_SYSTEM, "Konstruktionsmodell"),
    ("massnahmen", "Massnahmenteam", "Act",
     "Leitet je bestätigter Feststellung Ursachenhypothese, Sofort- und Korrekturmassnahme, "
     "Verantwortung, Frist und Wirksamkeitsnachweis ab.", MASSNAHMEN_SYSTEM,
     "Falsifikationsmodell")]

AGENT_STANDARD = {s: (n, z) for s, n, _, z, _, _ in AGENTEN}


def agent_name(audit: dict, schluessel: str) -> str:
    eigen = (audit.get("agenten") or {}).get(schluessel, {}).get("name", "").strip()
    return eigen or AGENT_STANDARD.get(schluessel, ("", ""))[0]


def agent_rolle(audit: dict, schluessel: str) -> str:
    return (audit.get("agenten") or {}).get(schluessel, {}).get("rolle", "")


URTEILE = ["Ausgeräumt", "Bestätigt: Abweichung", "Bestätigt: Verbesserungspotenzial",
           "Offen: vor Ort prüfen"]
TYPEN = {"widerlegend": "Nachweise sprechen dagegen",
         "untergrabend": "Nachweis ist nicht belastbar",
         "unterlaufend": "Nachweis deckt die Anforderung nicht ab",
         "ungestuetzte_annahme": "Behauptung ohne Nachweis"}
RISIKO = ["hoch", "mittel", "gering"]
STATUS = ["geplant", "in Vorbereitung", "durchgeführt", "berichtet", "abgeschlossen"]
MSTATUS = ["offen", "in Umsetzung", "umgesetzt", "wirksam bestätigt", "nicht wirksam"]


# ---------------------------------------------------------------------------
# Zugang
# ---------------------------------------------------------------------------
def anmeldung() -> bool:
    if st.session_state.get("auth_ok"):
        return True
    rand_l, mitte, rand_r = st.columns([1, 2, 1])
    with mitte:
        if LOGO.exists():
            bild_l, bild_m, bild_r = st.columns([1, 2, 1])
            with bild_m:
                st.image(str(LOGO), use_container_width=True)
        else:
            st.title(APP_NAME)
        st.markdown(
            f"<p style='text-align:center;font-size:1.1rem;color:#4a5a68;margin-top:-.5rem'>"
            f"{APP_CLAIM}</p>", unsafe_allow_html=True)
        st.write("")
        pw = st.text_input("Zugangspasswort", type="password",
                           placeholder="Passwort eingeben")
        if st.button("Anmelden", type="primary", use_container_width=True):
            if pw and pw == st.secrets.get("APP_PASSWORD", ""):
                st.session_state.auth_ok = True
                st.rerun()
            else:
                st.error("Das Passwort stimmt nicht.")
        with st.expander("Was Audix tut und was nicht"):
            st.markdown(HINWEISE)
    return False


if not anmeldung():
    st.stop()

st.session_state.setdefault("audits", [])


# ---------------------------------------------------------------------------
# Dateien und Normkataloge
# ---------------------------------------------------------------------------
@st.cache_data
def normkataloge_laden() -> dict:
    kataloge = {}
    if NORMEN_ORDNER.exists():
        for datei in sorted(NORMEN_ORDNER.glob("*.json")):
            try:
                k = json.loads(datei.read_text(encoding="utf-8"))
                kataloge[k.get("norm", datei.stem)] = k
            except (json.JSONDecodeError, OSError):
                pass
    return kataloge


def datei_lesen(hochgeladen) -> str:
    daten = hochgeladen.getvalue()
    name = hochgeladen.name.lower()
    if name.endswith(".pdf"):
        return "\n".join((s.extract_text() or "") for s in PdfReader(io.BytesIO(daten)).pages)
    if name.endswith(".docx"):
        doc = Document(io.BytesIO(daten))
        teile = [p.text for p in doc.paragraphs]
        for tab in doc.tables:
            for zeile in tab.rows:
                teile.append(" | ".join(z.text for z in zeile.cells))
        return "\n".join(teile)
    return daten.decode("utf-8", errors="ignore")


# ---------------------------------------------------------------------------
# Sprachmodell
# ---------------------------------------------------------------------------
class ModellFehler(Exception):
    pass


def json_aus_text(text: str):
    text = re.sub(r"^```(?:json)?|```$", "", text.strip(), flags=re.MULTILINE).strip()
    try:
        return json.loads(text)
    except json.JSONDecodeError:
        pass
    anfang, ende = text.find("{"), text.rfind("}")
    if anfang >= 0 and ende > anfang:
        try:
            return json.loads(text[anfang:ende + 1])
        except json.JSONDecodeError:
            pass
    return {"fehler": "Keine gültige JSON-Antwort", "rohtext": text[:2000]}


def llm_json(modell: str, system: str, nutzer: str, audit: dict, rolle: str = ""):
    """Anfrage an die OpenAI-kompatible Schnittstelle, bei Bedarf als Datenstrom."""
    name = agent_name(audit, rolle)
    eigene_rolle = agent_rolle(audit, rolle).strip()
    if name:
        system = f"Dein Name in diesem Audit ist {name}.\n" + system
    if eigene_rolle:
        system += ("\n\nROLLENVERSTÄNDNIS, das der Auditor für dich festgelegt hat:\n"
                   + eigene_rolle)
    zusatz = (audit.get("prompt_zusatz") or {}).get(rolle, "").strip()
    if zusatz:
        system += ("\n\nZUSÄTZLICHE ANWEISUNG DES AUDITORS. Sie ergänzt die Regeln oben und "
                   "hebt sie nicht auf:\n" + zusatz)
    system = system + NORMREGEL + SPRACHREGEL + JSONREGEL
    basis = {"model": modell, "temperature": 0.2, "max_tokens": MAX_TOKENS, "stream": STREAM,
             "response_format": {"type": "json_object"},
             "messages": [{"role": "system", "content": system},
                          {"role": "user", "content": nutzer}]}
    kopf = {"Authorization": f"Bearer {API_KEY}", "Content-Type": "application/json"}
    start = datetime.now()

    def senden(mit_format: bool):
        daten = dict(basis)
        if not mit_format:
            daten.pop("response_format", None)
        return requests.post(f"{API_BASE}/chat/completions", headers=kopf, json=daten,
                             timeout=(15, TIMEOUT), stream=STREAM)

    try:
        r = senden(True)
        if r.status_code == 400:
            r.close()
            r = senden(False)
    except requests.exceptions.Timeout:
        raise ModellFehler("Zeitüberschreitung. Weniger Nachweise je Lauf verwenden.")
    except requests.RequestException as f:
        raise ModellFehler(f"Verbindung fehlgeschlagen: {f}")

    if r.status_code != 200:
        if r.status_code in (502, 503, 504):
            raise ModellFehler(f"Der Dienst hat abgebrochen (HTTP {r.status_code}). Weniger "
                               "Nachweise laden oder MAX_TOKENS und MAX_ZEICHEN verkleinern.")
        raise ModellFehler(f"HTTP {r.status_code}: {r.text[:400]}")

    if STREAM:
        teile = []
        try:
            for zeile in r.iter_lines(decode_unicode=True):
                if not zeile or not zeile.startswith("data:"):
                    continue
                nutzdaten = zeile[5:].strip()
                if nutzdaten == "[DONE]":
                    break
                try:
                    stueck = json.loads(nutzdaten)
                except json.JSONDecodeError:
                    continue
                delta = (stueck.get("choices") or [{}])[0].get("delta", {})
                teile.append(delta.get("content") or "")
        except requests.RequestException as f:
            raise ModellFehler(f"Die Antwort wurde unterbrochen: {f}")
        inhalt = "".join(teile).strip()
    else:
        inhalt = (r.json().get("choices", [{}])[0].get("message", {}).get("content") or "").strip()

    audit.setdefault("protokoll", []).append(
        {"zeit": start.isoformat(timespec="seconds"), "rolle": rolle, "modell": modell,
         "dauer_sekunden": round((datetime.now() - start).total_seconds()),
         "zeichen_eingabe": len(nutzer), "zeichen_ausgabe": len(inhalt),
         "zusatzanweisung": zusatz, "ausgabe": inhalt[:6000]})
    if not inhalt:
        raise ModellFehler("Das Modell hat nichts zurückgegeben. Bitte erneut versuchen.")
    return json_aus_text(inhalt)


# ---------------------------------------------------------------------------
# Nachweise und Pruefung
# ---------------------------------------------------------------------------
def normalisieren(text: str) -> str:
    return re.sub(r"\s+", " ", text or "").strip().lower()


def zitate_pruefen(nachweise, quellen):
    for n in nachweise or []:
        z = normalisieren(n.get("zitat", ""))
        n["verifiziert"] = bool(z) and z in normalisieren(quellen.get(n.get("dokument", ""), ""))
    return nachweise


def quellen_sammeln(audit: dict) -> dict:
    quellen = dict(audit.get("dokumente", {}))
    for i, n in enumerate(audit.get("notizen", []), start=1):
        quellen[f"{n['typ']} {i} ({n.get('quelle', 'ohne Angabe')}, {n.get('datum', '')})"] = n.get("text", "")
    return quellen


def korpus(quellen: dict) -> str:
    return "".join(f"\n=== NACHWEIS: {n} ===\n{i}\n" for n, i in quellen.items())[:MAX_ZEICHEN]


def kriterium_text(katalog: dict, abschnitt: dict) -> str:
    punkte = "\n".join(f"- {p}" for p in abschnitt.get("pruefpunkte", []))
    return (f"Norm: {katalog['norm']}, Abschnitt {abschnitt['nr']} {abschnitt['titel']}\n\n"
            f"Anforderung\n{abschnitt['anforderung']}\n\nPrüfpunkte\n{punkte}")


def plan_vorschlagen(audit: dict, katalog: dict):
    quellen = quellen_sammeln(audit)
    liste = "\n".join(f"- {a['nr']} {a['titel']}: {a['anforderung'][:160]}"
                      for a in katalog["abschnitte"])
    probe = "".join(f"\n=== {n} ===\n{t[:1500]}\n" for n, t in quellen.items())[:MAX_ZEICHEN]
    return llm_json(MODEL_PRO, SCOPING_SYSTEM,
                    f"Norm {katalog['norm']}\n\nNormabschnitte\n{liste}\n\n"
                    f"Auditierter Prozess: {audit.get('prozess', '')}\n\n"
                    f"Vorhandene Nachweise\n{probe}", audit, "vorbereitung")


def analyse_durchfuehren(audit, katalog, abschnitt, mit_erwiderung=True, melden=None):
    quellen = quellen_sammeln(audit)
    kriterium = kriterium_text(katalog, abschnitt)
    material = korpus(quellen)
    nr = abschnitt["nr"]
    if melden:
        melden(f"Konstruktionsteam prüft Abschnitt {nr}")
    fall = llm_json(MODEL_PRO, PRO_SYSTEM,
                    f"Auditkriterium\n{kriterium}\n\nNachweise\n{material}", audit, "pro")
    for t in fall.get("teilaussagen", []):
        zitate_pruefen(t.get("nachweise"), quellen)
    if melden:
        melden(f"Falsifikationsteam greift Abschnitt {nr} an")
    einw = llm_json(MODEL_CONTRA, CONTRA_SYSTEM,
                    f"Auditkriterium\n{kriterium}\n\nAssurance Case\n"
                    f"{json.dumps(fall, ensure_ascii=False)}\n\nNachweise\n{material}",
                    audit, "contra")
    praefix = nr.replace(".", "_")
    for x in einw.get("einwaende", []):
        x["id"] = f"{praefix}-{x.get('id', 'E')}"
        x["abschnitt"] = nr
        zitate_pruefen(x.get("nachweise"), quellen)
    erw = {"erwiderungen": []}
    if mit_erwiderung and einw.get("einwaende"):
        if melden:
            melden(f"Konstruktionsteam erwidert zu Abschnitt {nr}")
        erw = llm_json(MODEL_PRO, ERWIDERUNG_SYSTEM,
                       f"Einwände\n{json.dumps(einw, ensure_ascii=False)}\n\n"
                       f"Nachweise\n{material}", audit, "erwiderung")
        for x in erw.get("erwiderungen", []):
            zitate_pruefen(x.get("nachweise"), quellen)
    audit.setdefault("analysen", {})[nr] = {
        "titel": abschnitt["titel"], "anforderung": abschnitt["anforderung"],
        "pruefpunkte": abschnitt.get("pruefpunkte", []), "kriterium": kriterium,
        "fall": fall, "einwaende": einw, "erwiderungen": erw,
        "modelle": {"pro": MODEL_PRO, "contra": MODEL_CONTRA},
        "zeit": datetime.now().isoformat(timespec="seconds")}


def alle_einwaende(audit: dict) -> list:
    return [e for erg in audit.get("analysen", {}).values()
            for e in erg.get("einwaende", {}).get("einwaende", [])]


def ampel_teilaussage(t: dict, einwaende: list, urteile: dict):
    """Gibt (Zeichen, Kurztext) fuer eine Teilaussage zurueck."""
    belegt = any(n.get("verifiziert") for n in t.get("nachweise", []))
    offen = [e for e in einwaende if e.get("ziel") == t.get("id")
             and urteile.get(e.get("id"), {}).get("urteil") != "Ausgeräumt"]
    schwer = [e for e in offen if e.get("typ") in ("widerlegend", "untergrabend")
              or e.get("schwere") == "hoch"]
    if not belegt or schwer:
        return "🔴", "ohne belastbaren Nachweis"
    if offen or t.get("annahmen"):
        return "🟡", "Nachweis vorhanden, Zweifel offen"
    return "🟢", "belegt"


def ampel_abschnitt(erg: dict, urteile: dict):
    einwaende = erg.get("einwaende", {}).get("einwaende", [])
    zeichen = [ampel_teilaussage(t, einwaende, urteile)[0]
               for t in erg.get("fall", {}).get("teilaussagen", [])]
    if not zeichen:
        return "🔴", "Kein Nachweis erarbeitet"
    gruen, gelb, rot = zeichen.count("🟢"), zeichen.count("🟡"), zeichen.count("🔴")
    if rot:
        lage = "Mögliche Abweichung, mindestens ein Prüfpunkt ohne Nachweis"
    elif gelb:
        lage = "Nachweis unvollständig, offene Zweifel"
    else:
        lage = "Anforderung durchgängig belegt"
    return ("🔴" if rot else "🟡" if gelb else "🟢"), f"{lage} ({gruen} belegt, {gelb} mit Zweifel, {rot} ohne Nachweis)"


def kennzahlen(audit: dict) -> dict:
    nachweise = [n for erg in audit.get("analysen", {}).values()
                 for t in erg.get("fall", {}).get("teilaussagen", [])
                 for n in t.get("nachweise", [])]
    nachweise += [n for e in alle_einwaende(audit) for n in e.get("nachweise", [])]
    urteile = [u.get("urteil") for u in audit.get("urteile", {}).values()]
    mass = audit.get("massnahmen", [])
    return {"abschnitte": len(audit.get("analysen", {})),
            "einwaende": len(alle_einwaende(audit)),
            "beurteilt": sum(1 for u in urteile if u),
            "zitate": len(nachweise),
            "zitate_belegt": sum(1 for n in nachweise if n.get("verifiziert")),
            "abweichungen": sum(1 for u in urteile if u == "Bestätigt: Abweichung"),
            "potenziale": sum(1 for u in urteile if u == "Bestätigt: Verbesserungspotenzial"),
            "offen": sum(1 for u in urteile if u == "Offen: vor Ort prüfen"),
            "ausgeraeumt": sum(1 for u in urteile if u == "Ausgeräumt"),
            "massnahmen": len(mass),
            "massnahmen_offen": sum(1 for m in mass
                                    if m.get("status") in ("offen", "in Umsetzung"))}


def nachweise_anzeigen(nachweise, leer_hinweis="Kein Nachweis angegeben."):
    if not nachweise:
        st.caption(leer_hinweis)
    for n in nachweise or []:
        marke = ("🟢 Zitat im Dokument gefunden" if n.get("verifiziert")
                 else "🔴 Zitat NICHT im Dokument gefunden, nicht verwendbar")
        st.markdown(f"> {n.get('zitat', '')}  \n*Quelle · {n.get('dokument', '')}* · {marke}")


# ---------------------------------------------------------------------------
# Bericht
# ---------------------------------------------------------------------------
def bericht_bloecke(audit: dict) -> dict:
    k = kennzahlen(audit)
    kopf = [("Organisation", ORGANISATION or "nicht angegeben"),
            ("Auditierter Bereich oder Prozess", audit.get("prozess", "")),
            ("Auditziel", audit.get("ziel", "")),
            ("Umfang und Grenzen", audit.get("umfang", "")),
            ("Auditkriterien", f"{audit.get('norm', '')}, Abschnitte "
                               + ", ".join(audit.get("kriterien", []))),
            ("Auditor", audit.get("auditor", "")),
            ("Unparteilichkeit bestätigt", "ja" if audit.get("unparteilich") else "nicht bestätigt"),
            ("Auditierte", audit.get("auditierte", "")),
            ("Audittermin", audit.get("termin", "")),
            ("Berichtsdatum", f"{date.today():%d.%m.%Y}")]
    quellen = ([f"Dokument · {n}" for n in audit.get("dokumente", {})]
               + [f"{n['typ']} · {n.get('quelle', '')} · {n.get('datum', '')}"
                  for n in audit.get("notizen", [])])
    return {"kennzahlen": k, "kopf": kopf, "quellen": quellen,
            "feststellungen": audit.get("feststellungen", {}).get("feststellungen", []),
            "zusammenfassung": audit.get("feststellungen", {}).get("zusammenfassung", ""),
            "fazit": audit.get("fazit", ""),
            "ausgeraeumt": [(i, u["begruendung"]) for i, u in audit.get("urteile", {}).items()
                            if u.get("urteil") == "Ausgeräumt"],
            "massnahmen": audit.get("massnahmen", [])}


def bericht_markdown(audit: dict) -> str:
    b = bericht_bloecke(audit)
    k = b["kennzahlen"]
    z = [f"# Auditbericht · {audit.get('titel', '')}", ""]
    z += [f"**{name}** {wert}  " for name, wert in b["kopf"]]
    z += ["", "## Zusammenfassung", b["zusammenfassung"] or "nicht erstellt", ""]
    if b["fazit"]:
        z += ["### Fazit des Auditors", b["fazit"], ""]
    z += ["## Ergebnis je Normabschnitt", "", "| Abschnitt | Ergebnis |", "| --- | --- |"]
    for nr, erg in audit.get("analysen", {}).items():
        zeichen, text = ampel_abschnitt(erg, audit.get("urteile", {}))
        z.append(f"| {nr} {erg.get('titel', '')} | {zeichen} {text} |")
    z += ["", "## Kennzahlen der Prüfung", "", "| Grösse | Wert |", "| --- | --- |",
          f"| Geprüfte Normabschnitte | {k['abschnitte']} |",
          f"| Geprüfte Zweifel insgesamt | {k['einwaende']} |",
          f"| Davon ausgeräumt | {k['ausgeraeumt']} |",
          f"| Abweichungen | {k['abweichungen']} |",
          f"| Verbesserungspotenziale | {k['potenziale']} |",
          f"| Vor Ort nachzuprüfen | {k['offen']} |",
          f"| Zitate in der Quelle gefunden | {k['zitate_belegt']} von {k['zitate']} |", ""]
    z += ["## Geprüfte Nachweise"] + [f"- {q}" for q in b["quellen"]] + [""]
    z += ["## Feststellungen"]
    if not b["feststellungen"]:
        z.append("Keine Feststellungen.")
    for f in b["feststellungen"]:
        z += [f"### {f.get('einwand')} · {f.get('einstufung')}",
              f"**Normabschnitt** {f.get('abschnitt', '')}  ",
              f"**Anforderung** {f.get('anforderung')}  ",
              f"**Objektiver Nachweis** {f.get('objektiver_nachweis')}  ",
              f"**Feststellung** {f.get('feststellung')}", ""]
    z += ["## Massnahmenplan"]
    if not b["massnahmen"]:
        z.append("Keine Massnahmen erfasst.")
    else:
        z += ["", "| Feststellung | Korrekturmassnahme | Verantwortlich | Termin | Status |",
              "| --- | --- | --- | --- | --- |"]
        for m in b["massnahmen"]:
            z.append(f"| {m.get('feststellung', '')} | {m.get('beschreibung', '')} | "
                     f"{m.get('verantwortlich', '')} | {m.get('termin', '')} | "
                     f"{m.get('status', '')} |")
        z.append("")
        for m in b["massnahmen"]:
            if m.get("ursache_hypothese") or m.get("wirksamkeitsnachweis"):
                z += [f"**{m.get('feststellung', '')}**  ",
                      f"Ursachenhypothese · {m.get('ursache_hypothese', '')}  ",
                      f"Sofortmassnahme · {m.get('sofortmassnahme', '')}  ",
                      f"Wirksamkeitsnachweis · {m.get('wirksamkeitsnachweis', '')}", ""]
    z += ["## Ausgeräumte Zweifel (Nachweis der Prüftiefe)"]
    z += [f"- {i}: {g}" for i, g in b["ausgeraeumt"]] or ["Keine."]
    z += ["", "---", "", f"Erstellt mit {APP_NAME}, einem dialektischen Prüfverfahren. Zwei getrennte "
          "Agententeams haben den Konformitätsnachweis konstruiert und angegriffen. Über jeden "
          "verbliebenen Zweifel hat der Auditor entschieden und dies begründet. Einstufung und "
          "Schlussfolgerung verantwortet der Auditor."]
    return "\n".join(z)


def bericht_html(audit: dict) -> str:
    b = bericht_bloecke(audit)
    k = b["kennzahlen"]

    def schutz(t):
        return str(t).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")

    kopf = "".join(f"<tr><th>{schutz(n)}</th><td>{schutz(w)}</td></tr>" for n, w in b["kopf"])
    kacheln = "".join(
        f'<div class="kachel"><div class="zahl">{w}</div><div class="text">{t}</div></div>'
        for t, w in [("Normabschnitte", k["abschnitte"]), ("Geprüfte Zweifel", k["einwaende"]),
                     ("Abweichungen", k["abweichungen"]), ("Potenziale", k["potenziale"]),
                     ("Vor Ort offen", k["offen"]),
                     ("Zitate belegt", f"{k['zitate_belegt']}/{k['zitate']}")])
    klasse = {"🟢": "gruen", "🟡": "gelb", "🔴": "rot"}
    uebersicht = "".join(
        f'<tr><td>{schutz(nr)} {schutz(erg.get("titel", ""))}</td>'
        f'<td class="{klasse[ampel_abschnitt(erg, audit.get("urteile", {}))[0]]}">'
        f'{schutz(ampel_abschnitt(erg, audit.get("urteile", {}))[1])}</td></tr>'
        for nr, erg in audit.get("analysen", {}).items())
    farbe = {"Bestätigt: Abweichung": "abw", "Bestätigt: Verbesserungspotenzial": "pot",
             "Offen: vor Ort prüfen": "off"}
    fest = "".join(
        f'<div class="fest {farbe.get(f.get("einstufung", ""), "")}">'
        f'<h3>{schutz(f.get("einwand"))} · {schutz(f.get("einstufung"))}</h3>'
        f'<p><b>Normabschnitt</b> {schutz(f.get("abschnitt", ""))}</p>'
        f'<p><b>Anforderung</b> {schutz(f.get("anforderung"))}</p>'
        f'<p><b>Objektiver Nachweis</b> {schutz(f.get("objektiver_nachweis"))}</p>'
        f'<p><b>Feststellung</b> {schutz(f.get("feststellung"))}</p></div>'
        for f in b["feststellungen"]) or "<p>Keine Feststellungen.</p>"
    mass = "".join(
        f"<tr><td>{schutz(m.get('feststellung', ''))}</td><td>{schutz(m.get('beschreibung', ''))}"
        f"<br><small>Ursachenhypothese: {schutz(m.get('ursache_hypothese', ''))}<br>"
        f"Wirksamkeitsnachweis: {schutz(m.get('wirksamkeitsnachweis', ''))}</small></td>"
        f"<td>{schutz(m.get('verantwortlich', ''))}</td><td>{schutz(m.get('termin', ''))}</td>"
        f"<td>{schutz(m.get('status', ''))}</td></tr>" for m in b["massnahmen"])
    quellen = "".join(f"<li>{schutz(q)}</li>" for q in b["quellen"])
    ausger = "".join(f"<li><b>{schutz(i)}</b> {schutz(g)}</li>" for i, g in b["ausgeraeumt"])
    fazit = f"<h2>Fazit des Auditors</h2><p>{schutz(b['fazit'])}</p>" if b["fazit"] else ""
    return f"""<!doctype html><html lang="de"><head><meta charset="utf-8">
<title>Auditbericht {schutz(audit.get('titel', ''))}</title><style>
body{{font-family:Segoe UI,Helvetica,Arial,sans-serif;max-width:900px;margin:2rem auto;
padding:0 1.5rem;color:#1a1a1a;line-height:1.55}}
h1{{font-size:1.7rem;border-bottom:3px solid #1f4e79;padding-bottom:.4rem;color:#1f4e79}}
h2{{font-size:1.2rem;margin-top:2rem;color:#1f4e79}}
table{{border-collapse:collapse;width:100%;margin:.6rem 0}}
th,td{{border:1px solid #d5dde5;padding:.45rem .6rem;text-align:left;vertical-align:top;
font-size:.92rem}}
th{{background:#f0f4f8;width:16rem}}
td.gruen{{background:#e8f5e9}} td.gelb{{background:#fff8e1}} td.rot{{background:#fdecea}}
.kacheln{{display:flex;flex-wrap:wrap;gap:.6rem;margin:1rem 0}}
.kachel{{flex:1 1 8rem;background:#f0f4f8;border-radius:8px;padding:.7rem;text-align:center}}
.zahl{{font-size:1.6rem;font-weight:700;color:#1f4e79}}
.text{{font-size:.78rem;color:#4a5a68}}
.fest{{border-left:5px solid #99a;background:#fafbfc;padding:.7rem 1rem;margin:.8rem 0;
border-radius:0 6px 6px 0}}
.fest.abw{{border-color:#b3261e}} .fest.pot{{border-color:#e08b00}} .fest.off{{border-color:#5a6b7a}}
.fest h3{{margin:.2rem 0 .5rem;font-size:1rem}} .fest p{{margin:.25rem 0;font-size:.92rem}}
small{{color:#55626e}} footer{{margin-top:2.5rem;font-size:.82rem;color:#55626e;
border-top:1px solid #d5dde5;padding-top:.8rem}}
@media print{{body{{margin:0}} .fest{{break-inside:avoid}}}}
</style></head><body>
<h1>Auditbericht · {schutz(audit.get('titel', ''))}</h1>
<table>{kopf}</table>
<h2>Zusammenfassung</h2><p>{schutz(b['zusammenfassung']) or 'nicht erstellt'}</p>
{fazit}
<h2>Ergebnis je Normabschnitt</h2>
<table><tr><th>Abschnitt</th><th>Ergebnis der Nachweisprüfung</th></tr>{uebersicht}</table>
<h2>Kennzahlen der Prüfung</h2><div class="kacheln">{kacheln}</div>
<h2>Feststellungen</h2>{fest}
<h2>Massnahmenplan</h2>
<table><tr><th>Feststellung</th><th>Massnahme</th><th>Verantwortlich</th><th>Termin</th>
<th>Status</th></tr>{mass or '<tr><td colspan="5">Keine Massnahmen erfasst.</td></tr>'}</table>
<h2>Geprüfte Nachweise</h2><ul>{quellen}</ul>
<h2>Ausgeräumte Zweifel</h2><ul>{ausger or '<li>Keine.</li>'}</ul>
<footer>Erstellt mit {schutz(APP_NAME)}, einem dialektischen Prüfverfahren. Zwei getrennte Agententeams haben den
Konformitätsnachweis konstruiert und angegriffen. Über jeden verbliebenen Zweifel hat der
Auditor entschieden und dies begründet. Einstufung und Schlussfolgerung verantwortet der
Auditor.</footer></body></html>"""


# ---------------------------------------------------------------------------
# Berichtsformular nach dem Muster eines Auditkurzberichts, auf ISO 9001 bezogen
# Aufbau (schluessel, Beschriftung, Hilfetext, Zeilen im Eingabefeld)
# ---------------------------------------------------------------------------
BERICHT_FELDER = [
    ("ziel_audit", "1.1 Zielsetzung des Audits",
     "Wozu wird auditiert. Wird aus dem Auditplan vorbelegt.", 100),
    ("dienstleistung", "1.2 Auditierter Prozess oder Bereich",
     "Welcher Prozess, welche Tätigkeiten und welche Schnittstellen.", 80),
    ("institution", "1.3 Organisation und Ansprechpartner",
     "Organisationseinheit, Anzahl Mitarbeitende, Prozessverantwortliche.", 80),
    ("zeitplan", "1.4 Auditzeitplan",
     "Datum, Uhrzeit, Ablauf, vor Ort oder remote.", 70),
    ("standort", "1.5 Geprüfter Standort",
     "Nur bei mehreren Standorten oder virtuellen Standorten ausfüllen.", 60),
    ("unterlagen", "1.6 Eingereichte Unterlagen",
     "Wird aus den erfassten Nachweisen vorbelegt.", 120),
    ("haftung", "1.7 Haftungsausschluss",
     "Standardtext. Die Prüfung beruht auf einer Stichprobe.", 90),
    ("verteiler", "1.8 Verteiler",
     "Wer den Bericht erhält. ISO 9001 verlangt Bericht an die zuständige Leitung.", 80),
    ("geltungsbereich", "2 Geltungsbereich",
     "Welcher Teil des QM-Systems geprüft wurde.", 80),
    ("begriffe", "3 Begriffe",
     "Definition von Abweichung, Hinweis und Empfehlung.", 110),
    ("veraenderungen_org", "4.1 Wichtige Veränderungen in der Organisation",
     "Organisatorische und personelle Änderungen seit dem letzten Audit.", 80),
    ("veraenderungen_prozess", "4.2 Veränderungen am Prozess oder Angebot",
     "Neue oder geänderte Prozesse, Produkte oder Dienstleistungen.", 80),
    ("kennzahlen_entwicklung", "4.3 Entwicklung der relevanten Kennzahlen",
     "Mengen, Reklamationen, Termintreue, Ausschuss und ähnliche Daten.", 80),
    ("erledigung_vorjahr", "4.4 Erledigungsnachweise früherer Korrekturmassnahmen",
     "Stand der Massnahmen aus früheren Audits und Bewertung ihrer Wirksamkeit.", 90),
    ("umgang_hinweise", "4.5 Umgang mit Hinweisen aus früheren Berichten",
     "Welche Hinweise wurden aufgegriffen, welche bleiben offen.", 80),
    ("qualitaetsinitiativen", "4.6 Eigene Qualitätsinitiativen",
     "Projekte, Evaluationen und Verbesserungen aus eigener Initiative.", 80),
    ("selbstbewertung", "4.7 Selbstbewertung der Organisation",
     "Wie die Organisation ihren eigenen Erfüllungsgrad einschätzt.", 80),
    ("regelkreis_fuehrung", "5.1.1 Regelkreis Führung und Qualitätsmanagement",
     "Normabschnitte 4, 5, 6, 9.3 und 10. Wird aus den Befunden vorbelegt.", 110),
    ("regelkreis_leistung", "5.1.2 Regelkreis Leistungserbringung und Unterstützung",
     "Normabschnitte 7 und 8. Wird aus den Befunden vorbelegt.", 110),
    ("gespraeche", "5.2 Eröffnungs- und Abschlussgespräch",
     "Wird aus den erfassten Gesprächsnotizen vorbelegt.", 100),
    ("erfuellung", "6.1 Erfüllung der geprüften Normabschnitte",
     "Wird aus der Ampelübersicht vorbelegt.", 120),
    ("begruendung", "6.2 Begründung des Auditergebnisses",
     "Warum das Ergebnis so ausfällt. Das ist Ihre Kernaussage als Auditor.", 120),
    ("hinweise_text", "6.3.3 Hinweise",
     "Entwicklungspotenziale und Risiken ohne Abweichungscharakter. "
     "Wird aus den Verbesserungspotenzialen vorbelegt.", 110),
    ("empfehlungen_text", "6.3.4 Empfehlungen",
     "Freiwillige Anregungen mit Nutzen für die Organisation.", 100),
    ("schlusswort", "6.4 Schlusswort",
     "Abschliessende Würdigung und Hinweis auf die Grundlagen der Bewertung.", 120),
    ("naechste_pruefung", "7 Planung der nächsten Überprüfung",
     "Termin, Schwerpunkte und Standorte des nächsten Audits.", 80),
    ("antrag", "8 Antrag an die Leitung",
     "Was der Auditor der Leitung empfiehlt, etwa Freigabe mit oder ohne Auflagen.", 100),
]


def feld_vorbelegung(audit: dict, schluessel: str) -> str:
    """Erzeugt den Vorschlagstext eines Berichtsfeldes aus den vorhandenen Daten."""
    k = kennzahlen(audit)
    urteile = audit.get("urteile", {})
    analysen = audit.get("analysen", {})

    if schluessel == "ziel_audit":
        basis = audit.get("ziel", "").strip()
        return (basis + "\n" if basis else "") + (
            "Das Audit verfolgt folgende Ziele:\n"
            f"- Prüfung der Konformität mit {audit.get('norm', 'der Norm')}, Abschnitte "
            + ", ".join(audit.get("kriterien", [])) + "\n"
            "- Beurteilung der Wirksamkeit des Prozesses und des Qualitätsmanagementsystems\n"
            "- Bewertung der Massnahmen aus früheren Audits\n"
            "- Ermittlung von Verbesserungspotenzialen\n"
            "- Berichterstattung an die zuständige Leitung")
    if schluessel == "dienstleistung":
        return audit.get("prozess", "")
    if schluessel == "institution":
        return (f"{ORGANISATION}\nAuditierte Funktion: {audit.get('auditierte', '')}\n"
                "Anzahl Mitarbeitende im geprüften Bereich: ")
    if schluessel == "zeitplan":
        return f"Audittermin: {audit.get('termin', '')}\nDurchführung: vor Ort"
    if schluessel == "standort":
        return "Nicht zutreffend, Einzelstandort."
    if schluessel == "unterlagen":
        zeilen = ["Folgende Nachweise wurden ausgewertet:"]
        zeilen += [f"- Dokument: {n}" for n in audit.get("dokumente", {})]
        zeilen += [f"- {n['typ']}: {n.get('quelle', '')} vom {n.get('datum', '')}"
                   for n in audit.get("notizen", [])]
        return "\n".join(zeilen)
    if schluessel == "haftung":
        return ("Das Audit beruht auf einem Stichprobenverfahren der verfügbaren Informationen. "
                f"Die Dokumentenanalyse wurde durch das KI-gestützte Prüfverfahren {APP_NAME} "
                "mit zwei "
                "getrennt arbeitenden Agententeams unterstützt. Jedes von der KI angeführte "
                "Zitat wurde maschinell gegen die Quelle geprüft. Die Bewertung der Nachweise "
                "und alle Schlussfolgerungen verantwortet der Auditor.")
    if schluessel == "verteiler":
        return ("Diesen Bericht erhalten:\n- die oberste Leitung\n"
                f"- die Leitung des auditierten Bereichs ({audit.get('auditierte', '')})\n"
                "- die Qualitätsmanagementbeauftragte Person\n"
                f"- der Auditor ({audit.get('auditor', '')})")
    if schluessel == "geltungsbereich":
        return (f"Geprüft wurde der Prozess {audit.get('prozess', '')} gegen "
                f"{audit.get('norm', '')}, Abschnitte "
                + ", ".join(audit.get("kriterien", [])) + ". "
                + (audit.get("umfang", "") or ""))
    if schluessel == "begriffe":
        return ("Die verwendeten Begriffe orientieren sich an ISO 9000 und ISO 19011.\n"
                "Abweichung (Nichtkonformität): Nichterfüllung einer Anforderung. "
                "Eine wesentliche Abweichung liegt vor, wenn eine Anforderung systematisch oder "
                "vollständig nicht erfüllt ist. Eine geringfügige Abweichung ist ein Einzelfall "
                "oder eine formale Lücke ohne Systemversagen.\n"
                "Hinweis: Entwicklungspotenzial oder Risiko ohne Abweichungscharakter.\n"
                "Empfehlung: freiwillige Anregung ohne Verpflichtung zur Umsetzung.")
    if schluessel in ("regelkreis_fuehrung", "regelkreis_leistung"):
        fuehrung = ("4", "5", "6", "9", "10")
        zeilen = []
        for nr, erg in analysen.items():
            ist_fuehrung = nr.split(".")[0] in fuehrung
            if (schluessel == "regelkreis_fuehrung") != ist_fuehrung:
                continue
            zeichen, lage = ampel_abschnitt(erg, urteile)
            zeilen.append(f"- {nr} {erg.get('titel', '')}: {lage}")
        return "\n".join(zeilen) or "In diesem Regelkreis wurde kein Abschnitt geprüft."
    if schluessel == "gespraeche":
        zeilen = [f"{n['typ']} am {n.get('datum', '')} mit {n.get('quelle', '')}\n{n.get('text', '')}"
                  for n in audit.get("notizen", [])
                  if n.get("typ") in ("Eröffnungsgespräch", "Abschlussgespräch")]
        return "\n\n".join(zeilen) or ("Eröffnungs- und Abschlussgespräch wurden durchgeführt. "
                                       "Inhalte bitte ergänzen.")
    if schluessel == "erfuellung":
        zeilen = []
        for nr, erg in analysen.items():
            zeichen, lage = ampel_abschnitt(erg, urteile)
            wort = {"🟢": "erfüllt", "🟡": "mit Einschränkung erfüllt",
                    "🔴": "nicht erfüllt"}[zeichen]
            zeilen.append(f"- {nr} {erg.get('titel', '')}: {wort}")
        zeilen.append("")
        zeilen.append(f"Geprüft wurden {k['abschnitte']} Normabschnitte. Das Prüfverfahren hat "
                      f"{k['einwaende']} begründete Zweifel erzeugt. Davon wurden "
                      f"{k['ausgeraeumt']} ausgeräumt, {k['abweichungen']} als Abweichung und "
                      f"{k['potenziale']} als Verbesserungspotenzial bestätigt, "
                      f"{k['offen']} bleiben zur Prüfung vor Ort offen.")
        return "\n".join(zeilen)
    if schluessel == "begruendung":
        if k["abweichungen"] == 0:
            return ("Die geprüften Anforderungen sind erfüllt. Für jede Anforderung lagen "
                    "Nachweise vor, deren Aussagekraft im Prüfverfahren gezielt angegriffen "
                    "wurde. Die aufgeworfenen Zweifel konnten mit den vorliegenden Nachweisen "
                    "ausgeräumt werden.")
        zahl = k["abweichungen"]
        wort = "eine Abweichung" if zahl == 1 else f"{zahl} Abweichungen"
        return (f"Es wurde {wort} festgestellt." if zahl == 1
                else f"Es wurden {wort} festgestellt.") + (
                " Die Begründung je Abweichung ist in Abschnitt 6.3 dargestellt. Bitte ergänzen "
                "Sie hier Ihre zusammenfassende Beurteilung der Systemreife, also ob es sich um "
                "Einzelfälle oder um eine systematische Schwäche handelt.")
    if schluessel == "hinweise_text":
        zeilen = ["Hinweise zeigen Entwicklungspotenziale und Risiken auf. Sie sind auf "
                  "Relevanz zu prüfen und dienen der fortlaufenden Verbesserung.", ""]
        for i, u in urteile.items():
            if u.get("urteil") == "Bestätigt: Verbesserungspotenzial":
                e = u.get("einwand", {})
                zeilen.append(f"- {i} (Abschnitt {e.get('abschnitt', '')}): "
                              f"{e.get('begruendung', '')[:400]}")
        if len(zeilen) == 2:
            zeilen.append("- Keine Hinweise.")
        return "\n".join(zeilen)
    if schluessel == "empfehlungen_text":
        return ("Empfehlungen sollen einen Nutzen für die Organisation stiften und stellen keine "
                "Abweichung dar. Die Umsetzung liegt im Ermessen der Organisation.\n"
                "- Bitte ergänzen.")
    if schluessel == "schlusswort":
        offen = [f"{i}: {u.get('einwand', {}).get('pruefung_vor_ort', '')}"
                 for i, u in urteile.items() if u.get("urteil") == "Offen: vor Ort prüfen"]
        text = [audit.get("fazit", "").strip(),
                f"Die Bewertung beruht auf {audit.get('norm', 'der Norm')} und auf den Leitlinien "
                "für Managementsystemaudits nach ISO 19011.",
                f"Das Audit wurde durch das dialektische KI-Prüfverfahren {APP_NAME} unterstützt. "
                "Ein Agententeam hat den Konformitätsnachweis konstruiert, ein zweites hat ihn "
                "angegriffen. Über jeden verbliebenen Zweifel hat der Auditor entschieden und "
                "diese Entscheidung begründet."]
        if offen:
            text.append("Folgende Punkte konnten anhand der Unterlagen nicht abschliessend "
                        "beurteilt werden und sind vor Ort zu prüfen:")
            text += [f"- {o}" for o in offen]
        return "\n".join(t for t in text if t)
    if schluessel == "veraenderungen_org":
        return ("Bitte ergänzen: organisatorische und personelle Veränderungen seit dem letzten "
                "Audit, etwa neue Funktionen, Wechsel in der Prozessverantwortung, "
                "Umstrukturierungen. Bei unveränderter Lage: Die Organisation ist unverändert.")
    if schluessel == "veraenderungen_prozess":
        return ("Bitte ergänzen: neue oder geänderte Prozesse, Produkte, Dienstleistungen, "
                "Anlagen oder IT-Systeme im geprüften Bereich. Bei unveränderter Lage: "
                "Der Prozess ist unverändert.")
    if schluessel == "kennzahlen_entwicklung":
        zahlen = [n for n in audit.get("notizen", []) if n.get("typ") == "Leistungsdaten"]
        if zahlen:
            return "\n".join(f"{n.get('quelle', '')} ({n.get('datum', '')}): {n.get('text', '')}"
                              for n in zahlen)
        return ("Bitte ergänzen: Entwicklung der relevanten Kennzahlen, etwa Mengen, "
                "Reklamationen, Termintreue, Ausschuss oder Durchlaufzeiten, mit Vergleich zur "
                "Vorperiode. Leistungsdaten können im Reiter Durchführung erfasst werden, dann "
                "erscheinen sie hier automatisch.")
    if schluessel == "erledigung_vorjahr":
        return ("Bitte ergänzen: Stand der Korrekturmassnahmen aus früheren Audits und Bewertung "
                "ihrer Wirksamkeit. Bei erstmaligem Audit oder ohne Vorbefunde: Im letzten Audit "
                "wurden keine Abweichungen festgestellt.")
    if schluessel == "umgang_hinweise":
        return ("Bitte ergänzen: welche Hinweise aus früheren Berichten aufgegriffen wurden und "
                "welche offen bleiben. Bei erstmaligem Audit: Nicht zutreffend.")
    if schluessel == "qualitaetsinitiativen":
        return ("Bitte ergänzen: Projekte, Evaluationen und Verbesserungen, die die Organisation "
                "aus eigener Initiative angestossen hat, einschliesslich der daraus abgeleiteten "
                "Massnahmen.")
    if schluessel == "selbstbewertung":
        return ("Bitte ergänzen: wie die Organisation ihren eigenen Erfüllungsgrad einschätzt und "
                "ob diese Selbsteinschätzung mit dem Auditbefund übereinstimmt. Eine Abweichung "
                "zwischen Selbstbild und Befund ist selbst ein Auditergebnis.")
    if schluessel == "naechste_pruefung":
        return ("Nächstes internes Audit dieses Prozesses: noch festzulegen.\n"
                "Schwerpunkte: Wirksamkeit der vereinbarten Korrekturmassnahmen"
                + (" sowie die oben offen gebliebenen Punkte." if k["offen"] else "."))
    if schluessel == "antrag":
        if k["abweichungen"] == 0:
            return ("Der Auditor empfiehlt der Leitung, die Konformität des geprüften Prozesses "
                    "zu bestätigen. Auflagen sind nicht erforderlich. Die Hinweise sollten im "
                    "Rahmen der fortlaufenden Verbesserung aufgegriffen werden.")
        return ("Der Auditor empfiehlt der Leitung, die festgestellten Abweichungen mit "
                "Korrekturmassnahmen zu belegen und deren Wirksamkeit innerhalb der "
                "vereinbarten Fristen nachzuweisen. Der Bericht ist in der nächsten "
                "Managementbewertung zu behandeln.")
    return ""


def felder_fuellen(audit: dict, nur_leere: bool = True):
    """Belegt die Berichtsfelder vor. Vom Auditor geaenderte Texte bleiben erhalten."""
    felder = audit.setdefault("bericht_felder", {})
    for schluessel, _, _, _ in BERICHT_FELDER:
        if nur_leere and felder.get(schluessel, "").strip():
            continue
        felder[schluessel] = feld_vorbelegung(audit, schluessel)
    return felder


def abweichungen_sortiert(audit: dict):
    """Teilt die bestaetigten Abweichungen in wesentliche und geringfuegige."""
    fest = {f.get("einwand"): f for f in
            audit.get("feststellungen", {}).get("feststellungen", [])}
    mass = {}
    for m in audit.get("massnahmen", []):
        mass.setdefault(m.get("feststellung"), m)
    major, minor = [], []
    for i, u in audit.get("urteile", {}).items():
        if u.get("urteil") != "Bestätigt: Abweichung":
            continue
        e = u.get("einwand", {})
        r = e.get("zertifizierungsrisiko") or {}
        klass = str(r.get("klassifizierung", "")).lower()
        schwer = ("haupt" in klass or "major" in klass or e.get("schwere") == "hoch")
        eintrag = {"id": i, "abschnitt": e.get("abschnitt", ""),
                   "feststellung": fest.get(i, {}).get("feststellung", e.get("begruendung", "")),
                   "anforderung": fest.get(i, {}).get("anforderung", ""),
                   "nachweis": fest.get(i, {}).get("objektiver_nachweis", ""),
                   "urteil_begruendung": u.get("begruendung", ""),
                   "massnahme": mass.get(i, {}).get("beschreibung", ""),
                   "sofort": mass.get(i, {}).get("sofortmassnahme", ""),
                   "verantwortlich": mass.get(i, {}).get("verantwortlich", ""),
                   "termin": mass.get(i, {}).get("termin", ""),
                   "wirksamkeit": mass.get(i, {}).get("wirksamkeitsnachweis", "")}
        (major if schwer else minor).append(eintrag)
    return major, minor


# ---------------------------------------------------------------------------
# Word-Bericht nach dem Schema eines Auditkurzberichts
# ---------------------------------------------------------------------------
def bericht_docx(audit: dict) -> bytes:
    from docx import Document as NeuesDok
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, RGBColor

    f = audit.get("bericht_felder", {})
    k = kennzahlen(audit)
    urteile = audit.get("urteile", {})
    major, minor = abweichungen_sortiert(audit)
    d = NeuesDok()

    stil = d.styles["Normal"]
    stil.font.name = "Calibri"
    stil.font.size = Pt(10.5)

    def absatz(text="", stilname=None, fett=False, kursiv=False, groesse=None):
        p = d.add_paragraph(style=stilname) if stilname else d.add_paragraph()
        for i, zeile in enumerate(str(text).split("\n")):
            lauf = p.add_run(("\n" if i else "") + zeile)
            lauf.bold = fett
            lauf.italic = kursiv
            if groesse:
                lauf.font.size = Pt(groesse)
        return p

    def feldtext(schluessel, ersatz="Nicht ausgefüllt."):
        wert = str(f.get(schluessel, "")).strip()
        if not wert:
            absatz(ersatz, kursiv=True)
            return
        for zeile in wert.split("\n"):
            z = zeile.strip()
            if not z:
                continue
            if z.startswith("- "):
                d.add_paragraph(z[2:], style="List Bullet")
            else:
                absatz(z)

    def tabelle(zeilen, breit=None):
        t = d.add_table(rows=0, cols=2)
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.LEFT
        for name, wert in zeilen:
            r = t.add_row().cells
            r[0].text = str(name)
            for p in r[0].paragraphs:
                for lauf in p.runs:
                    lauf.bold = True
            r[1].text = str(wert)
        d.add_paragraph()
        return t

    # ---------------- Titelseite
    titel = absatz(f"Bericht zum internen Audit", groesse=20, fett=True)
    titel.alignment = WD_ALIGN_PARAGRAPH.LEFT
    absatz(f"{audit.get('norm', '')}", groesse=14, fett=True)
    absatz()
    absatz(audit.get("titel", ""), groesse=14, fett=True)
    absatz(ORGANISATION or "", groesse=12)
    absatz()
    tabelle([("Audit", audit.get("titel", "")),
             ("Auditierter Prozess oder Bereich", audit.get("prozess", "")),
             ("Normative Grundlage", audit.get("norm", "")),
             ("Geprüfte Normabschnitte", ", ".join(audit.get("kriterien", []))),
             ("Bedeutung und Risiko des Prozesses", audit.get("risiko", "")),
             ("Audittermin", audit.get("termin", "")),
             ("Auditor", audit.get("auditor", "")),
             ("Unparteilichkeit bestätigt",
              "ja, der Auditor prüft nicht die eigene Arbeit"
              if audit.get("unparteilich") else "nicht bestätigt"),
             ("Auditierte Funktion", audit.get("auditierte", "")),
             ("Berichtsdatum", f"{date.today():%d.%m.%Y}"),
             ("Auditkennung", audit.get("id", ""))])
    absatz("Verfasser", fett=True)
    absatz(audit.get("auditor", ""))
    d.add_page_break()

    # Inhaltsverzeichnis als Word-Feld. Word fuellt es beim Oeffnen mit F9.
    d.add_heading("Inhaltsverzeichnis", level=1)
    p = d.add_paragraph()
    lauf = p.add_run()
    anfang = OxmlElement("w:fldChar")
    anfang.set(qn("w:fldCharType"), "begin")
    anweisung = OxmlElement("w:instrText")
    anweisung.set(qn("xml:space"), "preserve")
    anweisung.text = r'TOC \o "1-3" \h \z \u'
    trenner = OxmlElement("w:fldChar")
    trenner.set(qn("w:fldCharType"), "separate")
    platzhalter = OxmlElement("w:t")
    platzhalter.text = ("Inhaltsverzeichnis. In Word mit der rechten Maustaste darauf klicken "
                        "und Felder aktualisieren wählen.")
    schluss = OxmlElement("w:fldChar")
    schluss.set(qn("w:fldCharType"), "end")
    for teil in (anfang, anweisung, trenner, platzhalter, schluss):
        lauf._r.append(teil)
    d.add_page_break()

    # ---------------- 1 Allgemeine Informationen
    d.add_heading("1  Allgemeine Informationen", level=1)
    d.add_heading("1.1  Zielsetzung des Audits", level=2)
    feldtext("ziel_audit")
    d.add_heading("1.2  Auditierter Prozess oder Bereich", level=2)
    feldtext("dienstleistung")
    d.add_heading("1.3  Organisation und Ansprechpartner", level=2)
    feldtext("institution")
    d.add_heading("1.4  Auditzeitplan", level=2)
    feldtext("zeitplan")
    d.add_heading("1.5  Geprüfter Standort", level=2)
    feldtext("standort")
    d.add_heading("1.6  Eingereichte Unterlagen", level=2)
    feldtext("unterlagen")
    d.add_heading("1.7  Haftungsausschluss", level=2)
    feldtext("haftung")
    d.add_heading("1.8  Verteiler", level=2)
    feldtext("verteiler")

    d.add_heading("2  Geltungsbereich", level=1)
    feldtext("geltungsbereich")
    d.add_heading("3  Begriffe", level=1)
    feldtext("begriffe")

    # ---------------- 4 Auswertung
    d.add_heading("4  Auswertung der Organisation und des Prozesses", level=1)
    for schluessel, nummer, name in [
            ("veraenderungen_org", "4.1", "Wichtige Veränderungen in der Organisation"),
            ("veraenderungen_prozess", "4.2", "Veränderungen am Prozess oder Angebot"),
            ("kennzahlen_entwicklung", "4.3", "Entwicklung der relevanten Kennzahlen"),
            ("erledigung_vorjahr", "4.4", "Erledigungsnachweise früherer Korrekturmassnahmen"),
            ("umgang_hinweise", "4.5", "Umgang mit Hinweisen aus früheren Berichten"),
            ("qualitaetsinitiativen", "4.6", "Eigene Qualitätsinitiativen"),
            ("selbstbewertung", "4.7", "Selbstbewertung der Organisation")]:
        d.add_heading(f"{nummer}  {name}", level=2)
        feldtext(schluessel)

    # ---------------- 5 Auditprozess
    d.add_heading("5  Zusammenfassung des Auditprozesses und der geprüften Inhalte", level=1)
    d.add_heading("5.1  Ausführungen", level=2)
    d.add_heading("5.1.1  Regelkreis Führung und Qualitätsmanagement", level=3)
    feldtext("regelkreis_fuehrung")
    d.add_heading("5.1.2  Regelkreis Leistungserbringung und Unterstützung", level=3)
    feldtext("regelkreis_leistung")
    d.add_heading("5.1.3  Angewandtes Prüfverfahren", level=3)
    absatz("Die Dokumentenanalyse wurde durch ein Mehragentensystem unterstützt. "
           f"{agent_name(audit, 'pro')} hat je Normabschnitt den Konformitätsnachweis "
           f"konstruiert, {agent_name(audit, 'contra')} hat diesen Nachweis angegriffen und "
           "nach widersprechenden Nachweisen und ungestützten Annahmen gesucht. Jedes Zitat "
           "wurde maschinell gegen die Quelldatei geprüft. Über jeden verbliebenen Zweifel hat "
           "der Auditor entschieden und die Entscheidung begründet.")
    tabelle([("Geprüfte Normabschnitte", k["abschnitte"]),
             ("Erzeugte Zweifel", k["einwaende"]),
             ("Davon durch den Auditor ausgeräumt", k["ausgeraeumt"]),
             ("Als Abweichung bestätigt", k["abweichungen"]),
             ("Als Verbesserungspotenzial bestätigt", k["potenziale"]),
             ("Zur Prüfung vor Ort offen", k["offen"]),
             ("Zitate in der Quelle wiedergefunden",
              f"{k['zitate_belegt']} von {k['zitate']}")])
    d.add_heading("5.2  Eröffnungs- und Abschlussgespräch", level=2)
    feldtext("gespraeche")

    # ---------------- 6 Auditergebnis
    d.add_heading("6  Auditergebnis", level=1)
    d.add_heading("6.1  Erfüllung der geprüften Normabschnitte", level=2)
    t = d.add_table(rows=1, cols=3)
    t.style = "Table Grid"
    kopf = t.rows[0].cells
    for i, name in enumerate(["Normabschnitt", "Ergebnis", "Befund"]):
        kopf[i].text = name
        for p in kopf[i].paragraphs:
            for lauf in p.runs:
                lauf.bold = True
    for nr, erg in audit.get("analysen", {}).items():
        zeichen, lage = ampel_abschnitt(erg, urteile)
        wort = {"🟢": "erfüllt", "🟡": "mit Einschränkung erfüllt",
                "🔴": "nicht erfüllt"}[zeichen]
        r = t.add_row().cells
        r[0].text = f"{nr} {erg.get('titel', '')}"
        r[1].text = wort
        r[2].text = lage
    d.add_paragraph()
    feldtext("erfuellung", "")
    d.add_heading("6.2  Begründung", level=2)
    feldtext("begruendung")

    d.add_heading("6.3  Abweichungen und Korrekturmassnahmen", level=2)
    absatz("Eine Abweichung bezeichnet eine im Audit festgestellte Nichterfüllung einer "
           "Anforderung und ist mit einer Korrekturmassnahme zu belegen. Unterschieden werden "
           "wesentliche und geringfügige Abweichungen.")

    def abweichungsblock(titel, nummer, liste, frist_hinweis, leertext):
        d.add_heading(f"{nummer}  {titel}", level=3)
        if not liste:
            absatz(leertext)
            return
        for a in liste:
            absatz(f"{a['id']} · Normabschnitt {a['abschnitt']}", fett=True)
            tabelle([("Anforderung", a["anforderung"]),
                     ("Objektiver Nachweis", a["nachweis"]),
                     ("Feststellung", a["feststellung"]),
                     ("Begründung des Auditors", a["urteil_begruendung"]),
                     ("Sofortmassnahme", a["sofort"]),
                     ("Korrekturmassnahme", a["massnahme"]),
                     ("Verantwortlich", a["verantwortlich"]),
                     ("Termin", a["termin"]),
                     ("Nachweis der Wirksamkeit", a["wirksamkeit"])])
        absatz(frist_hinweis, kursiv=True)

    absatz("Eine wesentliche Abweichung liegt vor, wenn eine Anforderung systematisch oder "
           "vollständig nicht erfüllt ist. Eine geringfügige Abweichung ist ein Einzelfall oder "
           "eine formale Lücke ohne Systemversagen.", kursiv=True, groesse=9)
    abweichungsblock("Wesentliche Abweichungen mit Korrekturmassnahmen", "6.3.1", major,
                     "Nachweise zu den Korrekturmassnahmen sind termingerecht vorzulegen. "
                     "Die Wirksamkeit wird nachverfolgt.",
                     "Es wurden keine wesentlichen Abweichungen festgestellt.")
    abweichungsblock("Geringfügige Abweichungen mit Korrekturmassnahmen", "6.3.2", minor,
                     "Die Erledigung wird in der Regel beim nächsten Audit geprüft.",
                     "Es wurden keine geringfügigen Abweichungen festgestellt.")
    d.add_heading("6.3.3  Hinweise", level=3)
    feldtext("hinweise_text")
    d.add_heading("6.3.4  Empfehlungen", level=3)
    feldtext("empfehlungen_text")
    d.add_heading("6.4  Schlusswort", level=2)
    feldtext("schlusswort")

    # ---------------- 7 Planung
    d.add_heading("7  Planung der nächsten Überprüfung", level=1)
    feldtext("naechste_pruefung")

    # ---------------- 8 Antrag und Unterschriften
    d.add_heading("8  Antrag an die Leitung und Freigabe", level=1)
    feldtext("antrag")
    absatz()
    tabelle([("Auditor", f"{audit.get('auditor', '')}, Datum: ______________"),
             ("Unterschrift Auditor", "______________________________"),
             ("Zur Kenntnis genommen, Leitung des Bereichs", "______________________________"),
             ("Zur Kenntnis genommen, oberste Leitung", "______________________________")])
    absatz("Dieser Bericht ist als dokumentierte Information aufzubewahren "
           "(ISO 9001, Abschnitte 9.2 und 7.5).", kursiv=True, groesse=9)

    # ---------------- Anhang
    d.add_page_break()
    d.add_heading("Anhang A  Ausgeräumte Zweifel als Nachweis der Prüftiefe", level=1)
    absatz("Diese Zweifel wurden im Prüfverfahren erzeugt und vom Auditor mit Begründung "
           "ausgeräumt. Sie dokumentieren die Tiefe der Prüfung.")
    ausger = [(i, u) for i, u in urteile.items() if u.get("urteil") == "Ausgeräumt"]
    if not ausger:
        absatz("Keine.")
    for i, u in ausger:
        e = u.get("einwand", {})
        absatz(f"{i} · Normabschnitt {e.get('abschnitt', '')}", fett=True)
        absatz(f"Zweifel: {e.get('begruendung', '')}")
        absatz(f"Begründung der Ausräumung: {u.get('begruendung', '')}")
        absatz()

    d.add_heading("Anhang B  Zur Prüfung vor Ort offene Punkte", level=1)
    offen = [(i, u) for i, u in urteile.items() if u.get("urteil") == "Offen: vor Ort prüfen"]
    if not offen:
        absatz("Keine.")
    for i, u in offen:
        e = u.get("einwand", {})
        v = e.get("pruefung_vor_ort")
        v = v if isinstance(v, list) else ([v] if v else [])
        absatz(f"{i} · Normabschnitt {e.get('abschnitt', '')}", fett=True)
        absatz(f"Offener Zweifel: {e.get('begruendung', '')}")
        for schritt in v:
            d.add_paragraph(str(schritt), style="List Bullet")
        absatz()

    puffer = io.BytesIO()
    d.save(puffer)
    return puffer.getvalue()


# ---------------------------------------------------------------------------
# Oberflaeche
# ---------------------------------------------------------------------------
kataloge = normkataloge_laden()
audits = st.session_state.audits

if not st.session_state.get("begruesst"):
    audix_begruessung()

kopf_text, kopf_hilfe = st.columns([5, 1])
with kopf_text:
    st.markdown(f"## {APP_NAME}")
    st.caption(((ORGANISATION + " · ") if ORGANISATION else "")
               + "Internes Audit nach ISO 9001 entlang des PDCA-Zyklus")
st.divider()

with st.sidebar:
    bereich = st.radio("Bereich", ["Plan · Auditprogramm", "Do und Check · Einzelaudit",
                                   "Act · Nachverfolgung"], label_visibility="collapsed")
    st.divider()
    st.subheader("Arbeitsstand")
    st.caption("Nichts wird dauerhaft gespeichert. Stand als Datei sichern und später laden.")
    st.download_button("Arbeitsstand sichern", json.dumps(audits, ensure_ascii=False, indent=1),
                       file_name=f"auditstand_{date.today()}.json", use_container_width=True)
    wieder = st.file_uploader("Arbeitsstand laden", type=["json"], key="restore")
    if wieder is not None and st.button("Übernehmen", use_container_width=True):
        try:
            st.session_state.audits = json.loads(wieder.getvalue().decode("utf-8"))
            st.rerun()
        except (json.JSONDecodeError, UnicodeDecodeError):
            st.error("Die Datei konnte nicht gelesen werden.")
    st.divider()
    if st.button("Einführung von Audix", use_container_width=True):
        st.session_state.begruesst = False
        st.rerun()
    with st.expander("Was Audix tut"):
        st.markdown(HINWEISE)
    with st.expander("Technik"):
        st.caption(f"Konstruktionsmodell {MODEL_PRO}\n\nFalsifikationsmodell {MODEL_CONTRA}")
        schnell = st.checkbox("Schnellmodus ohne Erwiderungsrunde", value=False)
        if not API_KEY:
            st.error("Kein API-Schlüssel in den Secrets hinterlegt.")

# ================================================================ Plan
if bereich.startswith("Plan"):
    st.session_state["aktuelles_audit"] = None
    with kopf_hilfe:
        audix_hilfe("programm")
    st.subheader("Auditprogramm des Jahres")
    if not kataloge:
        st.error("Kein Normkatalog gefunden. Legen Sie eine JSON-Datei im Ordner normen ab.")
        st.stop()
    with st.form("neu"):
        s1, s2, s3 = st.columns(3)
        titel = s1.text_input("Bezeichnung", "Prozessaudit Einkauf")
        prozess = s2.text_input("Prozess oder Bereich", "Einkauf")
        risiko = s3.selectbox("Bedeutung und Risiko", RISIKO)
        auditor = s1.text_input("Auditor")
        auditierte = s2.text_input("Auditierte Funktion")
        termin = s3.date_input("Geplanter Termin", date.today())
        norm = st.selectbox("Norm", list(kataloge))
        liste = {f"{a['nr']} {a['titel']}": a["nr"] for a in kataloge[norm]["abschnitte"]}
        wahl = st.multiselect("Zu prüfende Normabschnitte (leer lassen, wenn das System "
                              "sie anhand der Nachweise vorschlagen soll)", list(liste))
        unparteilich = st.checkbox("Der Auditor prüft nicht die eigene Arbeit "
                                   "(Unparteilichkeit nach ISO 9001 Abschnitt 9.2)")
        if st.form_submit_button("Audit ins Programm aufnehmen", type="primary"):
            if not unparteilich:
                st.error("Ohne bestätigte Unparteilichkeit darf das Audit nicht geplant werden.")
            else:
                audits.append({"id": date.today().strftime("%Y%m%d") + "-" + uuid.uuid4().hex[:4],
                               "titel": titel, "prozess": prozess, "risiko": risiko,
                               "termin": str(termin), "auditor": auditor,
                               "auditierte": auditierte, "norm": norm,
                               "kriterien": [liste[w] for w in wahl], "ziel": "", "umfang": "",
                               "unparteilich": True, "status": "geplant", "dokumente": {},
                               "notizen": [], "analysen": {}, "urteile": {},
                               "feststellungen": {}, "massnahmen": [], "protokoll": [],
                               "prompt_zusatz": {}, "fazit": ""})
                st.success("Audit angelegt. Weiter im Bereich Einzelaudit.")
                st.rerun()

    if audits:
        st.markdown("**Jahresübersicht**")
        st.dataframe([{"Bezeichnung": a["titel"], "Prozess": a.get("prozess", ""),
                       "Risiko": a.get("risiko", ""), "Termin": a.get("termin", ""),
                       "Auditor": a.get("auditor", ""),
                       "Abschnitte": ", ".join(a.get("kriterien", [])) or "noch offen",
                       "Abweichungen": kennzahlen(a)["abweichungen"],
                       "Offene Massnahmen": kennzahlen(a)["massnahmen_offen"],
                       "Status": a.get("status", "")} for a in audits],
                     use_container_width=True)
        st.markdown("**Abdeckung der Normabschnitte über das Programm**")
        norm_wahl = st.selectbox("Norm", list(kataloge), key="abd")
        alle = [a["nr"] for a in kataloge[norm_wahl]["abschnitte"]]
        geplant = {nr for a in audits if a.get("norm") == norm_wahl
                   for nr in a.get("kriterien", [])}
        offen = [nr for nr in alle if nr not in geplant]
        s1, s2 = st.columns(2)
        s1.metric("Abgedeckte Abschnitte", f"{len(alle) - len(offen)} / {len(alle)}")
        s2.metric("Noch offen", len(offen))
        if offen:
            st.info("Noch nicht im Programm enthalten: " + ", ".join(offen))
        else:
            st.success("Alle Abschnitte sind im Programm abgedeckt.")
    else:
        st.info("Noch keine Audits im Programm.")

# ================================================================ Do und Check
elif bereich.startswith("Do"):
    if not audits:
        st.warning("Legen Sie zuerst im Auditprogramm ein Audit an.")
        st.stop()
    namen = {f"{a['termin']} · {a['titel']}": a for a in audits}
    with st.sidebar:
        st.subheader("Audit")
        audit = namen[st.selectbox("Audit", list(namen), label_visibility="collapsed")]
    st.session_state["aktuelles_audit"] = audit
    katalog = kataloge.get(audit.get("norm"))
    if not katalog:
        st.error("Der Normkatalog dieses Audits wurde nicht gefunden.")
        st.stop()
    abschnitte = {a["nr"]: a for a in katalog["abschnitte"]}
    k = kennzahlen(audit)
    audit.setdefault("prompt_zusatz", {})

    with st.sidebar:
        st.divider()
        st.subheader("Fortschritt")
        schritte = [("Vorbereitung", bool(audit.get("kriterien"))),
                    ("Nachweise gesammelt", bool(quellen_sammeln(audit))),
                    ("Prüfung gelaufen", bool(audit.get("analysen"))),
                    ("Urteile vollständig", k["einwaende"] > 0 and k["beurteilt"] == k["einwaende"]),
                    ("Bericht erstellt", bool(audit.get("feststellungen"))),
                    ("Massnahmen festgelegt", bool(audit.get("massnahmen")))]
        st.progress(sum(1 for _, f in schritte if f) / len(schritte))
        for name, fertig in schritte:
            st.markdown(f"{name} · {'erledigt' if fertig else 'offen'}")

    t1, t2, t3, t4, t5, t6 = st.tabs(
        ["1 Vorbereitung (Plan)", "2 Durchführung (Do)", "3 Prüfung (Check)",
         "4 Urteil (Check)", "5 Bericht (Check)", "6 Massnahmen (Act)"])

    # ---------------- 1 Vorbereitung (Plan)
    with t1:
        hilfe_l, hilfe_r = st.columns([5, 1])
        with hilfe_r:
            audix_hilfe("vorbereitung")
        with st.form("plan"):
            s1, s2 = st.columns(2)
            audit["titel"] = s1.text_input("Bezeichnung", audit.get("titel", ""))
            audit["prozess"] = s2.text_input("Prozess oder Bereich", audit.get("prozess", ""))
            audit["auditor"] = s1.text_input("Auditor", audit.get("auditor", ""))
            audit["auditierte"] = s2.text_input("Auditierte", audit.get("auditierte", ""))
            audit["termin"] = str(s1.date_input(
                "Termin", date.fromisoformat(audit.get("termin", str(date.today())))))
            audit["status"] = s2.selectbox("Status", STATUS,
                                           index=STATUS.index(audit.get("status", "geplant"))
                                           if audit.get("status") in STATUS else 0)
            audit["ziel"] = st.text_area("Auditziel", audit.get("ziel", ""),
                                         placeholder="Feststellung der Konformität und "
                                                     "Wirksamkeit des Prozesses")
            audit["umfang"] = st.text_area("Umfang und Grenzen", audit.get("umfang", ""),
                                           placeholder="Standort, Zeitraum, betrachtete "
                                                       "Tätigkeiten, Stichprobenumfang")
            liste = {f"{a['nr']} {a['titel']}": a["nr"] for a in katalog["abschnitte"]}
            vorauswahl = [t for t, v in liste.items() if v in audit.get("kriterien", [])]
            neu = st.multiselect("Zu prüfende Normabschnitte", list(liste), default=vorauswahl)
            audit["unparteilich"] = st.checkbox(
                "Unparteilichkeit bestätigt, der Auditor prüft nicht die eigene Arbeit",
                value=bool(audit.get("unparteilich")))
            if st.form_submit_button("Auditplan speichern", type="primary"):
                audit["kriterien"] = [liste[n] for n in neu]
                st.success("Gespeichert.")
        if not audit.get("unparteilich"):
            st.warning("Die Unparteilichkeit ist nicht bestätigt. ISO 9001 Abschnitt 9.2 "
                       "verlangt, dass Auditoren ihre eigene Arbeit nicht auditieren.")

        st.divider()
        st.markdown("**Was in den gewählten Abschnitten geprüft wird**")
        st.caption("Diese Übersicht ist zugleich Ihre Auditcheckliste und die Liste der "
                   "Unterlagen, die Sie beim Fachbereich anfordern sollten.")
        if not audit.get("kriterien"):
            st.info("Noch keine Normabschnitte gewählt. Entweder oben auswählen oder im Reiter "
                    "Prüfung vom System vorschlagen lassen.")
        for nr in audit.get("kriterien", []):
            a = abschnitte.get(nr)
            if not a:
                continue
            with st.expander(f"{nr} {a['titel']}", expanded=False):
                st.markdown(f"**Was die Norm verlangt**  \n{a['anforderung']}")
                st.markdown("**Worauf der Auditor achtet**")
                for p in a.get("pruefpunkte", []):
                    st.markdown(f"- {p}")
                if a.get("typische_nachweise"):
                    st.markdown("**Diese Unterlagen sollten Sie bereitstellen**")
                    for n in a["typische_nachweise"]:
                        st.markdown(f"- {n}")
        st.divider()
        st.markdown("**Auditteam zusammenstellen**")
        st.caption("Jede Rolle ist ein eigener Agent und lässt sich umbenennen, im "
                   "Rollenverständnis schärfen und mit einer Zusatzanweisung versehen. Das "
                   "dialektische Prinzip bleibt erhalten, denn ein Team baut den Nachweis und "
                   "ein zweites greift ihn an.")
        audit.setdefault("agenten", {})
        for schluessel, standardname, phase, zweck, prompt, modell in AGENTEN:
            eintrag = audit["agenten"].setdefault(schluessel, {"name": "", "rolle": ""})
            with st.container(border=True):
                st.markdown(f"**{standardname}** · Phase {phase} · {modell}")
                st.caption(zweck)
                s1, s2 = st.columns([1, 2])
                eintrag["name"] = s1.text_input(
                    "Name dieses Teams", value=eintrag.get("name", ""),
                    key=f"an_{audit['id']}_{schluessel}", placeholder=standardname)
                eintrag["rolle"] = s2.text_area(
                    "Rollenverständnis", value=eintrag.get("rolle", ""), height=70,
                    key=f"ar_{audit['id']}_{schluessel}",
                    placeholder="Beispiel: Du bist ein erfahrener Auditor aus der Medizintechnik "
                                "und legst besonderen Wert auf Rückverfolgbarkeit.")
                audit["prompt_zusatz"][schluessel] = st.text_area(
                    "Zusätzliche Anweisung für dieses Audit",
                    value=audit["prompt_zusatz"].get(schluessel, ""), height=70,
                    key=f"pz_{audit['id']}_{schluessel}",
                    placeholder="Beispiel: Achte besonders auf Fristen und auf Dokumente ohne "
                                "Freigabevermerk.")
                with st.expander("Grundanweisung dieses Teams anzeigen"):
                    st.code(prompt, language="text")

        if audit.get("kriterien"):
            zeilen = [f"# Auditplan {audit['titel']}", "",
                      f"Termin {audit['termin']}", f"Prozess {audit['prozess']}",
                      f"Auditor {audit['auditor']}", f"Auditierte {audit['auditierte']}",
                      "", f"## Auditziel\n{audit.get('ziel', '')}",
                      f"## Umfang\n{audit.get('umfang', '')}",
                      f"## Auditkriterien\n{audit['norm']}, Abschnitte "
                      + ", ".join(audit["kriterien"]), "", "## Checkliste"]
            for nr in audit["kriterien"]:
                a = abschnitte.get(nr, {})
                zeilen += [f"### {nr} {a.get('titel', '')}", a.get("anforderung", ""), "",
                           "Prüfpunkte"]
                zeilen += [f"- {p}" for p in a.get("pruefpunkte", [])]
                zeilen += ["", "Bereitzustellende Unterlagen"]
                zeilen += [f"- {n}" for n in a.get("typische_nachweise", [])] + [""]
            st.download_button("Auditplan mit Checkliste herunterladen", "\n".join(zeilen),
                               file_name=f"auditplan_{audit['id']}.md")

    # ---------------- 2 Durchfuehrung (Do)
    with t2:
        hilfe_l, hilfe_r = st.columns([5, 1])
        with hilfe_r:
            audix_hilfe("durchfuehrung")
        neue = st.file_uploader("Dokumente hinzufügen (PDF, Word, Text)",
                                type=["pdf", "docx", "txt", "md"], accept_multiple_files=True,
                                key=f"up_{audit['id']}")
        if neue and st.button("Dokumente übernehmen", type="primary"):
            for f in neue:
                audit.setdefault("dokumente", {})[f.name] = datei_lesen(f)
            st.rerun()
        for name in list(audit.get("dokumente", {})):
            s1, s2 = st.columns([6, 1])
            s1.write(name)
            if s2.button("entfernen", key=f"del_{name}"):
                del audit["dokumente"][name]
                st.rerun()
        st.divider()
        st.markdown("**Interviews, Beobachtungen und Leistungsdaten**")
        st.caption("Halten Sie Aussagen möglichst wörtlich fest. Das Falsifikationsteam "
                   "vergleicht sie mit den Vorgabedokumenten und findet so Widersprüche "
                   "zwischen Vorschrift und gelebter Praxis.")
        with st.form("notiz", clear_on_submit=True):
            s1, s2, s3 = st.columns(3)
            typ = s1.selectbox("Art", ["Interviewnotiz", "Beobachtung", "Leistungsdaten",
                                       "Eröffnungsgespräch", "Abschlussgespräch"])
            quelle = s2.text_input("Quelle (Rolle, Arbeitsplatz, Kennzahl)")
            datum = s3.date_input("Datum", date.today())
            text = st.text_area("Notiz", height=120)
            if st.form_submit_button("Notiz speichern") and text.strip():
                audit.setdefault("notizen", []).append(
                    {"typ": typ, "quelle": quelle, "datum": str(datum), "text": text})
                st.rerun()
        for i, n in enumerate(audit.get("notizen", [])):
            with st.container(border=True):
                st.markdown(f"**{n['typ']}** · {n.get('quelle', '')} · {n.get('datum', '')}")
                st.write(n["text"])
                if st.button("löschen", key=f"nd_{i}"):
                    audit["notizen"].pop(i)
                    st.rerun()

    # ---------------- 3 Dialektische Pruefung (Check)
    with t3:
        hilfe_l, hilfe_r = st.columns([5, 1])
        with hilfe_r:
            audix_hilfe("pruefung")

        with st.expander("Welche Teams hier arbeiten"):
            for schluessel, standardname, phase, zweck, prompt, modell in AGENTEN:
                if phase != "Check":
                    continue
                st.markdown(f"**{agent_name(audit, schluessel)}** · {modell}")
                st.caption(zweck)
                if agent_rolle(audit, schluessel):
                    st.caption("Rollenverständnis · " + agent_rolle(audit, schluessel))
            st.caption("Namen, Rollen und Zusatzanweisungen stellen Sie im Reiter "
                       "Vorbereitung ein.")

        if not quellen_sammeln(audit):
            st.warning("Bitte zuerst im Reiter Durchführung Nachweise erfassen.")
        else:
            start = st.button("Prüfung starten", type="primary", use_container_width=True)
            if start:
                balken = st.progress(0.0, text="Vorbereitung")
                protokoll = st.empty()
                try:
                    schritte_gesamt = max(1, len(audit.get("kriterien") or [1]) * 3 + 1)
                    zaehler = {"n": 0}

                    def melden(text):
                        zaehler["n"] += 1
                        balken.progress(min(1.0, zaehler["n"] / schritte_gesamt), text=text)
                        protokoll.caption(text)

                    if not audit.get("kriterien"):
                        melden("Vorbereitungsteam wählt die Normabschnitte")
                        vor = plan_vorschlagen(audit, katalog)
                        gueltig = [n for n in vor.get("kriterien", []) if n in abschnitte]
                        audit["kriterien"] = gueltig or list(abschnitte)[:3]
                        audit["ziel"] = vor.get("ziel", audit.get("ziel", ""))
                        audit["umfang"] = vor.get("umfang", audit.get("umfang", ""))
                        audit["planbegruendung"] = vor.get("begruendungen", {})
                        audit["fehlende_nachweise"] = vor.get("fehlende_nachweise", [])
                        schritte_gesamt = len(audit["kriterien"]) * 3 + 1
                    for nr in audit["kriterien"]:
                        if nr in audit.get("analysen", {}):
                            continue
                        analyse_durchfuehren(audit, katalog, abschnitte[nr], not schnell, melden)
                    audit["status"] = "durchgeführt"
                    balken.progress(1.0, text="Prüfung abgeschlossen")
                    st.rerun()
                except ModellFehler as f:
                    st.error(str(f))

            if audit.get("analysen"):
                s1, s2, s3, s4 = st.columns(4)
                s1.metric("Geprüfte Abschnitte", k["abschnitte"])
                s2.metric("Gefundene Zweifel", k["einwaende"])
                s3.metric("Zitate belegt", f"{k['zitate_belegt']} / {k['zitate']}")
                s4.metric("Noch zu beurteilen", k["einwaende"] - k["beurteilt"])
                if k["zitate"] and k["zitate_belegt"] < k["zitate"]:
                    st.warning(f"{k['zitate'] - k['zitate_belegt']} Zitate wurden in Ihren "
                               "Unterlagen nicht wiedergefunden. Diese Nachweise sind nicht "
                               "verwendbar, sie sind rot markiert.")
            if audit.get("fehlende_nachweise"):
                st.info("Das Vorbereitungsteam vermisst: " + "; ".join(audit["fehlende_nachweise"]))

            with st.expander("Wie Sie die Ergebnisse lesen"):
                st.markdown(LEGENDE)

            urteile = audit.get("urteile", {})
            risiko_zeichen = {"hauptabweichung": "🔴", "nebenabweichung": "🟡",
                              "kein risiko": "🟢"}
            for nr, erg in audit.get("analysen", {}).items():
                einwaende_abschnitt = erg.get("einwaende", {}).get("einwaende", [])
                erwid = {x.get("einwand"): x
                         for x in erg.get("erwiderungen", {}).get("erwiderungen", [])}
                zeichen, lage = ampel_abschnitt(erg, urteile)
                with st.container(border=True):
                    st.markdown(f"## {zeichen} Abschnitt {nr} {erg.get('titel', '')}")
                    st.markdown("**1 Normforderung (Soll-Zustand)**")
                    st.write(erg.get("anforderung", ""))
                    st.markdown(f"**Gesamtbefund des Systems** · {zeichen} {lage}")
                    st.caption("Befund des Systems auf Basis der Unterlagen. Die Feststellung "
                               "trifft der Auditor im Reiter Urteil.")
                    for t in erg.get("fall", {}).get("teilaussagen", []):
                        z, kurz = ampel_teilaussage(t, einwaende_abschnitt, urteile)
                        zweifel = [e for e in einwaende_abschnitt if e.get("ziel") == t.get("id")]
                        titel = t.get("pruefpunkt") or t.get("aussage") or t.get("id")
                        with st.expander(f"{z} Prüfpunkt · {titel}", expanded=False):

                            st.markdown("**2 Audit-Befund und Status**")
                            meine = [urteile.get(e["id"], {}).get("urteil") for e in zweifel]
                            meine = [m for m in meine if m]
                            belegte = sum(1 for n in t.get("nachweise", [])
                                          if n.get("verifiziert"))
                            st.markdown(
                                f"- System-Befund · {z} {kurz}\n"
                                f"- Datenbasis · {belegte} von {len(t.get('nachweise', []))} "
                                f"Zitaten in den Unterlagen gefunden, {len(zweifel)} begründete "
                                f"Zweifel, {len(t.get('annahmen', []))} Behauptungen ohne Nachweis\n"
                                f"- Ihr Urteil · "
                                + (", ".join(sorted(set(meine))) if meine
                                   else "noch nicht entschieden"))

                            st.markdown("**3 Analyse und Argumentation (Ist-Zustand)**")
                            st.markdown(f"*Sichtweise {agent_name(audit, 'pro')}*")
                            st.write(t.get("aussage", ""))
                            if t.get("argument"):
                                st.caption(t["argument"])
                            st.markdown(f"*Sichtweise {agent_name(audit, 'contra')}*")
                            if zweifel:
                                for e in zweifel:
                                    st.write(e.get("begruendung", ""))
                                    kurz_e = (erwid.get(e["id"].split("-", 1)[-1])
                                              or erwid.get(e["id"]))
                                    if kurz_e:
                                        st.caption(
                                            f"Erwiderung {agent_name(audit, 'erwiderung')} · "
                                            + ("zugestanden · " if kurz_e.get("zugestanden")
                                               else "") + str(kurz_e.get("erwiderung", "")))
                            else:
                                st.write("Kein Widerspruch gegen diesen Prüfpunkt.")

                            st.markdown("**4 Dokumentenprüfung (die Beweise)**")
                            if not t.get("nachweise"):
                                st.markdown("- Kein Dokument eingereicht, das diesen Prüfpunkt "
                                            "belegt. Das ist der Grund für die rote Ampel.")
                            for n in t.get("nachweise", []):
                                if n.get("verifiziert"):
                                    st.markdown(
                                        f"- Dokument · **{n.get('dokument', '')}**  \n"
                                        f"  Status · 🟢 verwendbar  \n"
                                        f"  Zitat · {n.get('zitat', '')}")
                                else:
                                    st.markdown(
                                        f"- Dokument · **{n.get('dokument', '')}**  \n"
                                        f"  Status · 🔴 nicht verwendbar  \n"
                                        f"  Grund · Das zitierte Textstück wurde im Dokument "
                                        f"nicht gefunden. Es gibt dafür keinen schriftlichen "
                                        f"Nachweis.")
                            for a in t.get("annahmen", []):
                                st.markdown(f"- Behauptung ohne Nachweis · {a}")

                            st.markdown("**5 Audit-Zweifel (Risikoanalyse der Unterlagen)**")
                            if not zweifel:
                                st.markdown("- Keine.")
                            for e in zweifel:
                                schwere = {"hoch": "🔴 hoch", "mittel": "🟡 mittel",
                                           "gering": "🟢 gering"}.get(e.get("schwere"),
                                                                     e.get("schwere", ""))
                                st.markdown(
                                    f"- ID · `{e.get('id')}`  \n"
                                    f"  Art · {TYPEN.get(e.get('typ'), e.get('typ'))}  \n"
                                    f"  Schweregrad · {schwere}  \n"
                                    f"  Begründung · {e.get('begruendung', '')}")
                                nachweise_anzeigen(e.get("nachweise"),
                                                   "Kein Gegenbeleg. Der Zweifel stützt sich "
                                                   "auf das Fehlen eines Nachweises.")

                            st.markdown("**6 Empfehlung für die Prüfung vor Ort**")
                            schritte_vo = []
                            for e in zweifel:
                                v = e.get("pruefung_vor_ort")
                                schritte_vo += v if isinstance(v, list) else ([v] if v else [])
                            if schritte_vo:
                                for i, v in enumerate(schritte_vo, start=1):
                                    st.markdown(f"{i}. {v}")
                            else:
                                st.markdown("Keine besondere Prüfung vorgeschlagen.")

                            st.markdown("**7 Risiko im externen Zertifizierungsaudit**")
                            risiken = [e.get("zertifizierungsrisiko") for e in zweifel
                                       if isinstance(e.get("zertifizierungsrisiko"), dict)]
                            if not risiken:
                                st.caption("Das Falsifikationsteam hat dazu nichts geliefert. "
                                           "Schätzen Sie es im Urteil selbst ein.")
                            for r in risiken:
                                klass = str(r.get("klassifizierung", ""))
                                zr = risiko_zeichen.get(klass.lower(), "⚪")
                                st.markdown(
                                    f"- Klassifizierung · {zr} {klass}  \n"
                                    f"  Prognose · {r.get('prognose', '')}  \n"
                                    f"  Konsequenz · {r.get('konsequenz', '')}  \n"
                                    f"  Dringlichkeit · {r.get('dringlichkeit', '')}")
                            st.caption("Diese Einschätzung ist eine Prognose des Systems und "
                                       "bindet keine Zertifizierungsstelle.")
                    if st.button(f"Abschnitt {nr} erneut prüfen", key=f"rm_{nr}"):
                        del audit["analysen"][nr]
                        st.rerun()

            with st.expander("Agentenprotokoll, wer wann was geliefert hat"):
                if not audit.get("protokoll"):
                    st.caption("Noch kein Lauf.")
                for p in audit.get("protokoll", []):
                    st.markdown(f"**{p.get('rolle')}** · {p.get('modell')} · {p.get('zeit')} · "
                                f"{p.get('dauer_sekunden')} Sekunden · "
                                f"{p.get('zeichen_eingabe')} Zeichen Eingabe")
                    if p.get("zusatzanweisung"):
                        st.caption("Ergänzung des Auditors · " + p["zusatzanweisung"])
                    with st.expander("Rohausgabe"):
                        st.code(p.get("ausgabe", ""), language="json")

    # ---------------- 4 Urteil (Check)
    with t4:
        hilfe_l, hilfe_r = st.columns([5, 1])
        with hilfe_r:
            audix_hilfe("urteil")
        einwaende = alle_einwaende(audit)
        if not einwaende:
            st.info("Noch keine Zweifel. Führen Sie zuerst die Prüfung durch.")
        else:
            erw = {x.get("einwand"): x for erg in audit.get("analysen", {}).values()
                   for x in erg.get("erwiderungen", {}).get("erwiderungen", [])}
            st.progress(k["beurteilt"] / max(1, k["einwaende"]),
                        text=f"{k['beurteilt']} von {k['einwaende']} Zweifeln beurteilt")
            nur_offen = st.checkbox("Nur noch unbeurteilte anzeigen", value=False)
            for e in einwaende:
                eid = e["id"]
                alt = audit.get("urteile", {}).get(eid, {})
                if nur_offen and alt.get("urteil"):
                    continue
                with st.container(border=True):
                    st.markdown(f"#### {eid} · {TYPEN.get(e.get('typ'), e.get('typ'))} "
                                f"· Schwere {e.get('schwere')}")
                    st.caption(f"Normabschnitt {e.get('abschnitt')} · "
                               f"betrifft Prüfpunkt {e.get('ziel')}")
                    links, rechts = st.columns(2)
                    with links:
                        st.markdown("**Zweifel des Falsifikationsteams**")
                        st.write(e.get("begruendung"))
                        nachweise_anzeigen(e.get("nachweise"),
                                           "Kein Gegenbeleg, der Einwand stützt sich auf das "
                                           "Fehlen eines Nachweises.")
                    with rechts:
                        st.markdown("**Erwiderung des Konstruktionsteams**")
                        kurz = erw.get(eid.split("-", 1)[-1]) or erw.get(eid)
                        if kurz:
                            st.write(("zugestanden · " if kurz.get("zugestanden") else "")
                                     + str(kurz.get("erwiderung")))
                            nachweise_anzeigen(kurz.get("nachweise"), "Kein weiterer Nachweis.")
                        else:
                            st.write("keine Erwiderung")
                    v = e.get("pruefung_vor_ort")
                    v = v if isinstance(v, list) else ([v] if v else [])
                    if v:
                        st.caption("Vorschlag Prüfung vor Ort · " + " | ".join(str(x) for x in v))
                    r = e.get("zertifizierungsrisiko")
                    if isinstance(r, dict):
                        st.caption(f"Prognose externes Audit · {r.get('klassifizierung', '')} · "
                                   f"{r.get('konsequenz', '')}")
                    idx = URTEILE.index(alt["urteil"]) if alt.get("urteil") in URTEILE else None
                    urteil = st.radio("Ihr Urteil", URTEILE, key=f"u_{eid}", index=idx,
                                      horizontal=True)
                    begr = st.text_area("Begründung (Pflicht)", value=alt.get("begruendung", ""),
                                        key=f"b_{eid}", height=70)
                    audit.setdefault("urteile", {})[eid] = {"urteil": urteil,
                                                            "begruendung": begr, "einwand": e}
            fehlend = [i for i, u in audit.get("urteile", {}).items()
                       if not (u.get("urteil") and u.get("begruendung", "").strip())]
            if fehlend:
                st.warning("Ohne begründetes Urteil: " + ", ".join(fehlend))
            if st.button("Urteile festschreiben, Bericht und Massnahmen erzeugen",
                         type="primary", disabled=bool(fehlend), use_container_width=True):
                relevant = [u for u in audit["urteile"].values() if u["urteil"] != "Ausgeräumt"]
                balken = st.progress(0.0, text="Feststellungen werden formuliert")
                try:
                    audit["feststellungen"] = llm_json(
                        MODEL_PRO, BERICHT_SYSTEM,
                        f"Auditkriterien {audit['norm']}, Abschnitte "
                        f"{', '.join(audit['kriterien'])}\n\nEntscheidungen des Auditors\n"
                        f"{json.dumps(relevant, ensure_ascii=False)}", audit, "bericht")
                    fs = audit["feststellungen"].get("feststellungen", [])
                    balken.progress(0.5, text="Massnahmenteam leitet Verbesserungen ab")
                    if fs:
                        vor = llm_json(MODEL_CONTRA, MASSNAHMEN_SYSTEM,
                                       "Bestätigte Feststellungen\n"
                                       f"{json.dumps(fs, ensure_ascii=False)}", audit,
                                       "massnahmen")
                        vorhanden = {m.get("feststellung") for m in audit.get("massnahmen", [])}
                        for m in vor.get("massnahmen", []):
                            if m.get("feststellung") in vorhanden:
                                continue
                            try:
                                tage = int(m.get("frist_tage", 30))
                            except (TypeError, ValueError):
                                tage = 30
                            audit.setdefault("massnahmen", []).append({
                                "id": uuid.uuid4().hex[:6],
                                "feststellung": m.get("feststellung", ""),
                                "beschreibung": m.get("korrekturmassnahme", ""),
                                "sofortmassnahme": m.get("sofortmassnahme", ""),
                                "ursache_hypothese": m.get("ursache_hypothese", ""),
                                "wirksamkeitsnachweis": m.get("wirksamkeitsnachweis", ""),
                                "verantwortlich": m.get("verantwortlich_rolle", ""),
                                "termin": str(date.today() + timedelta(days=tage)),
                                "status": "offen", "wirksam_geprueft": "", "audit": audit["id"],
                                "audit_titel": audit.get("titel", "")})
                    audit["status"] = "berichtet"
                    balken.progress(1.0, text="Fertig")
                    st.rerun()
                except ModellFehler as f:
                    st.error(str(f))

    # ---------------- 5 Bericht (Check)
    with t5:
        hilfe_l, hilfe_r = st.columns([5, 1])
        with hilfe_r:
            audix_hilfe("bericht")
        if not audit.get("feststellungen"):
            st.info("Der Bericht entsteht, sobald Sie alle Zweifel beurteilt haben.")
        else:
            s1, s2, s3, s4 = st.columns(4)
            s1.metric("Abweichungen", k["abweichungen"])
            s2.metric("Potenziale", k["potenziale"])
            s3.metric("Vor Ort offen", k["offen"])
            s4.metric("Ausgeräumt", k["ausgeraeumt"])
            st.markdown("**Ergebnis je Normabschnitt**")
            st.dataframe([{"Abschnitt": f"{nr} {erg.get('titel', '')}",
                           "Ergebnis": " ".join(ampel_abschnitt(erg, audit.get("urteile", {})))}
                          for nr, erg in audit.get("analysen", {}).items()],
                         use_container_width=True, hide_index=True)
            audit["fazit"] = st.text_area(
                "Fazit des Auditors", value=audit.get("fazit", ""), height=100,
                help="Ihre Gesamteinschätzung. Sie fliesst in das Schlusswort des Berichts ein.")

            st.divider()
            st.markdown("### Berichtsformular")
            st.caption("Das Formular folgt dem Aufbau eines Auditberichts. Die Felder sind aus "
                       "den erfassten Daten vorbelegt und überschreibbar. Änderungen stehen "
                       "genau so im Word-Bericht.")
            audit.setdefault("bericht_felder", {})
            s1, s2 = st.columns(2)
            if s1.button("Leere Felder vorbelegen", use_container_width=True):
                felder_fuellen(audit, nur_leere=True)
                st.rerun()
            if s2.button("Alle Felder neu vorbelegen", use_container_width=True,
                         help="Überschreibt auch Ihre eigenen Texte."):
                felder_fuellen(audit, nur_leere=False)
                st.rerun()
            if not any(v.strip() for v in audit["bericht_felder"].values()):
                felder_fuellen(audit, nur_leere=True)

            gruppen = [("1  Allgemeine Informationen",
                        ["ziel_audit", "dienstleistung", "institution", "zeitplan", "standort",
                         "unterlagen", "haftung", "verteiler"]),
                       ("2 und 3  Geltungsbereich und Begriffe",
                        ["geltungsbereich", "begriffe"]),
                       ("4  Auswertung der Organisation und des Prozesses",
                        ["veraenderungen_org", "veraenderungen_prozess",
                         "kennzahlen_entwicklung", "erledigung_vorjahr", "umgang_hinweise",
                         "qualitaetsinitiativen", "selbstbewertung"]),
                       ("5  Auditprozess und geprüfte Inhalte",
                        ["regelkreis_fuehrung", "regelkreis_leistung", "gespraeche"]),
                       ("6  Auditergebnis",
                        ["erfuellung", "begruendung", "hinweise_text", "empfehlungen_text",
                         "schlusswort"]),
                       ("7 und 8  Planung und Antrag",
                        ["naechste_pruefung", "antrag"])]
            beschriftung = {sch: (lab, hilfe, hoehe)
                            for sch, lab, hilfe, hoehe in BERICHT_FELDER}
            for gruppentitel, schluessel_liste in gruppen:
                with st.expander(gruppentitel, expanded=False):
                    for sch in schluessel_liste:
                        lab, hilfe, hoehe = beschriftung[sch]
                        audit["bericht_felder"][sch] = st.text_area(
                            lab, value=audit["bericht_felder"].get(sch, ""), height=hoehe,
                            help=hilfe, key=f"bf_{audit['id']}_{sch}")

            major, minor = abweichungen_sortiert(audit)
            st.markdown("**Abweichungen im Bericht**")
            st.caption("Die Einteilung ergibt sich aus dem Urteil des Auditors und aus der "
                       "Prognose zum externen Audit. Wesentlich heisst, dass eine Anforderung "
                       "systematisch oder vollständig nicht erfüllt ist.")
            s1, s2 = st.columns(2)
            s1.metric("Wesentliche Abweichungen", len(major))
            s2.metric("Geringfügige Abweichungen", len(minor))
            if major or minor:
                st.dataframe(
                    [{"Art": art, "ID": a["id"], "Abschnitt": a["abschnitt"],
                      "Feststellung": a["feststellung"][:160],
                      "Korrekturmassnahme": a["massnahme"][:120],
                      "Verantwortlich": a["verantwortlich"], "Termin": a["termin"]}
                     for art, liste in [("wesentlich", major), ("geringfügig", minor)]
                     for a in liste],
                    use_container_width=True, hide_index=True)
            if any(not a["massnahme"] for a in major + minor):
                st.warning("Zu mindestens einer Abweichung fehlt noch eine Korrekturmassnahme. "
                           "Ergänzen Sie sie im Reiter Massnahmen, sonst bleibt die Tabelle "
                           "im Bericht leer.")

            st.divider()
            st.markdown("### Bericht herunterladen")
            s1, s2 = st.columns(2)
            s1.download_button(
                "Auditbericht als Word-Datei", bericht_docx(audit),
                file_name=f"Auditbericht_{audit.get('prozess', 'Prozess')}_{audit['id']}.docx",
                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                type="primary", use_container_width=True)
            s2.download_button("Bericht als Webseite (HTML)", bericht_html(audit),
                               file_name=f"auditbericht_{audit['id']}.html", mime="text/html",
                               use_container_width=True)
            s1.caption("Vollständiger Bericht nach dem Schema, in Word weiter bearbeitbar.")
            s2.caption("Kurzfassung zum Ansehen, über Drucken als PDF speicherbar.")
            s3, s4 = st.columns(2)
            s3.download_button("Bericht als Text (Markdown)", bericht_markdown(audit),
                               file_name=f"auditbericht_{audit['id']}.md",
                               use_container_width=True)
            trail = {kk: vv for kk, vv in audit.items() if kk != "dokumente"}
            trail["nachweise"] = list(audit.get("dokumente", {}))
            s4.download_button("Audit Trail (JSON)",
                               json.dumps(trail, ensure_ascii=False, indent=2),
                               file_name=f"audit_trail_{audit['id']}.json",
                               use_container_width=True)
            with st.expander("Kurzfassung ansehen"):
                st.markdown(bericht_markdown(audit))

    # ---------------- 6 Massnahmen (Act)
    with t6:
        hilfe_l, hilfe_r = st.columns([5, 1])
        with hilfe_r:
            audix_hilfe("massnahmen")
        if not audit.get("massnahmen"):
            st.info("Noch keine Massnahmen. Sie entstehen nach dem Festschreiben der Urteile.")
        else:
            st.warning("Die Ursachen sind Hypothesen. Sie gehören vor der Freigabe mit den "
                       "Prozessverantwortlichen geprüft.")
            tabelle = st.data_editor(
                audit["massnahmen"], num_rows="dynamic", use_container_width=True,
                key=f"me_{audit['id']}",
                column_config={
                    "id": None, "audit": None, "audit_titel": None,
                    "feststellung": st.column_config.TextColumn("Feststellung", width="small"),
                    "beschreibung": st.column_config.TextColumn("Korrekturmassnahme",
                                                                width="large"),
                    "sofortmassnahme": st.column_config.TextColumn("Sofortmassnahme"),
                    "ursache_hypothese": st.column_config.TextColumn("Ursachenhypothese"),
                    "wirksamkeitsnachweis": st.column_config.TextColumn("Wirksamkeitsnachweis"),
                    "verantwortlich": st.column_config.TextColumn("Verantwortlich"),
                    "termin": st.column_config.TextColumn("Termin", width="small"),
                    "status": st.column_config.SelectboxColumn("Status", options=MSTATUS),
                    "wirksam_geprueft": st.column_config.TextColumn("Wirksamkeit geprüft am")})
            if st.button("Massnahmen übernehmen", type="primary"):
                for zeile in tabelle:
                    zeile.setdefault("id", uuid.uuid4().hex[:6])
                    zeile.setdefault("audit", audit["id"])
                    zeile.setdefault("audit_titel", audit.get("titel", ""))
                    zeile.setdefault("status", "offen")
                    zeile.setdefault("wirksam_geprueft", "")
                audit["massnahmen"] = list(tabelle)
                st.success("Gespeichert. Die Nachverfolgung finden Sie links im Bereich Act.")

# ================================================================ Act
else:
    st.session_state["aktuelles_audit"] = None
    with kopf_hilfe:
        audix_hilfe("nachverfolgung")
    st.subheader("Nachverfolgung der Wirksamkeit")
    zeilen = [dict(m, Audit=a.get("titel", ""),
                   ueberfaellig="ja" if (m.get("status") in ("offen", "in Umsetzung")
                                         and str(m.get("termin", "")) < str(date.today())) else "")
              for a in audits for m in a.get("massnahmen", [])]
    if not zeilen:
        st.info("Noch keine Massnahmen. Sie entstehen im Einzelaudit nach Ihrem Urteil.")
    else:
        s1, s2, s3, s4 = st.columns(4)
        s1.metric("Massnahmen gesamt", len(zeilen))
        s2.metric("Offen", sum(1 for z in zeilen if z.get("status") in ("offen", "in Umsetzung")))
        s3.metric("Überfällig", sum(1 for z in zeilen if z["ueberfaellig"]))
        s4.metric("Wirksam bestätigt",
                  sum(1 for z in zeilen if z.get("status") == "wirksam bestätigt"))
        st.dataframe([{"Audit": z["Audit"], "Feststellung": z.get("feststellung", ""),
                       "Massnahme": z.get("beschreibung", ""),
                       "Verantwortlich": z.get("verantwortlich", ""),
                       "Termin": z.get("termin", ""), "Status": z.get("status", ""),
                       "Wirksamkeit geprüft": z.get("wirksam_geprueft", ""),
                       "überfällig": z["ueberfaellig"]} for z in zeilen],
                     use_container_width=True, hide_index=True)
        st.divider()
        st.markdown("**Wirksamkeit einer Massnahme prüfen**")
        auswahl = {f"{z['Audit']} · {z.get('feststellung', '')} · {z.get('beschreibung', '')[:50]}":
                   (z.get("audit"), z.get("id")) for z in zeilen}
        gewaehlt = st.selectbox("Massnahme", list(auswahl))
        aid, mid = auswahl[gewaehlt]
        ziel = next((m for a in audits if a["id"] == aid
                     for m in a.get("massnahmen", []) if m.get("id") == mid), None)
        if ziel:
            st.caption(f"Vereinbarter Wirksamkeitsnachweis · {ziel.get('wirksamkeitsnachweis', '')}")
            s1, s2 = st.columns(2)
            ergebnis = s1.selectbox("Ergebnis der Prüfung",
                                    ["wirksam bestätigt", "nicht wirksam", "umgesetzt"])
            pruefdatum = s2.date_input("Geprüft am", date.today())
            nachweis = st.text_area("Was wurde geprüft und mit welchem Ergebnis",
                                    value=ziel.get("wirksamkeit_notiz", ""), height=90)
            if st.button("Wirksamkeitsprüfung festhalten", type="primary"):
                ziel["status"] = ergebnis
                ziel["wirksam_geprueft"] = str(pruefdatum)
                ziel["wirksamkeit_notiz"] = nachweis
                st.success("Dokumentiert.")
                st.rerun()
